import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

class BrandTokens:
    BG_COLOR = RGBColor(3, 12, 30)         # Dark Blue #030C1E
    CARD_BG = RGBColor(8, 22, 50)          # Card Slate #081632
    CYAN = RGBColor(0, 255, 255)           # Aqua Background 2 #00FFFF
    GOLD = RGBColor(255, 184, 28)          # Gold Text 2 #FFB81C
    TEXT_WHITE = RGBColor(255, 255, 255)   # White #FFFFFF

    FONT_TITLE = "Times New Roman"
    FONT_BODY = "Arial"

class NativePPTXBuilder:
    def __init__(self):
        self.prs = pptx.Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def build_slide(self, slide_ast: dict) -> pptx.slide.Slide:
        blank_slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_slide_layout)

        # 1. Apply Slide Background Color
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BrandTokens.BG_COLOR

        # 2. Render Nodes at exact Spatial Coordinates
        for node in slide_ast.get("nodes", []):
            node_type = node.get("type")
            bbox = node.get("bbox", {"left": 1.0, "top": 1.0, "width": 4.0, "height": 2.0})

            if node_type in ("RECTANGLE_SHAPE", "CONTAINER_SHAPE"):
                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(bbox["left"]), Inches(bbox["top"]),
                    Inches(bbox["width"]), Inches(bbox["height"])
                )
                card.fill.solid()
                card.fill.fore_color.rgb = BrandTokens.CARD_BG
                card.line.color.rgb = BrandTokens.CYAN
                card.line.width = Pt(1.5)
                if hasattr(card, 'adjustments') and len(card.adjustments) > 0:
                    card.adjustments[0] = 0.04

            elif node_type == "OVAL_SHAPE":
                oval = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL,
                    Inches(bbox["left"]), Inches(bbox["top"]),
                    Inches(bbox["width"]), Inches(bbox["height"])
                )
                oval.fill.solid()
                oval.fill.fore_color.rgb = BrandTokens.CARD_BG
                oval.line.color.rgb = BrandTokens.CYAN

            elif node_type == "ARROW_SHAPE":
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(bbox["left"]), Inches(bbox["top"]),
                    Inches(bbox["width"]), Inches(bbox["height"])
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = BrandTokens.GOLD
                arrow.line.color.rgb = BrandTokens.GOLD

            elif node_type == "TEXT_CONTAINER":
                txBox = slide.shapes.add_textbox(
                    Inches(bbox["left"]), Inches(bbox["top"]),
                    Inches(bbox["width"]), Inches(bbox["height"])
                )
                tf = txBox.text_frame
                tf.word_wrap = True

                for i, para_data in enumerate(node.get("paragraphs", [])):
                    txt = para_data.get("text", "").strip()
                    if not txt:
                        continue

                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = txt

                    is_title = bbox["top"] < 1.5 and bbox["height"] < 1.2

                    if is_title:
                        p.font.name = BrandTokens.FONT_TITLE
                        p.font.size = Pt(24)
                        p.font.bold = True
                        p.font.color.rgb = BrandTokens.GOLD
                    else:
                        p.font.name = BrandTokens.FONT_BODY
                        p.font.size = Pt(para_data.get("font_size", 11))
                        p.font.bold = para_data.get("bold", False)
                        p.font.italic = para_data.get("italic", False)
                        p.font.underline = para_data.get("underline", False)
                        p.font.color.rgb = BrandTokens.TEXT_WHITE

        return slide

    def save(self, output_path: str):
        self.prs.save(output_path)
