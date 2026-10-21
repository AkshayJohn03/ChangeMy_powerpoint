import base64
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import PP_ALIGN
from typing import Dict, List, Any, Optional, Generator, Tuple

try:
    from brandmorph.engine.archetype_engine import SlideGraphClassifier, _is_footer_or_slidenum
except Exception:
    try:
        from .brandmorph.engine.archetype_engine import SlideGraphClassifier, _is_footer_or_slidenum
    except Exception:
        try:
            from backend.brandmorph.engine.archetype_engine import SlideGraphClassifier, _is_footer_or_slidenum
        except Exception:
            SlideGraphClassifier = None
            _is_footer_or_slidenum = None

# Minimum dimension in inches to avoid zero-division / negative-extent bugs
_MIN_EXTENT = 0.05

def _color_to_hex(color_obj) -> Optional[str]:
    if color_obj is None:
        return None
    try:
        if hasattr(color_obj, "rgb") and color_obj.rgb is not None:
            return f"#{str(color_obj.rgb).upper()}"
    except Exception:
        pass
    return None

class SlideDecompiler:
    def __init__(self, pptx_path: str, raw_mode: bool = False):
        self.raw_mode = raw_mode
        self.prs = pptx.Presentation(pptx_path)
        self.orig_width = self.prs.slide_width.inches if self.prs.slide_width else 13.333
        self.orig_height = self.prs.slide_height.inches if self.prs.slide_height else 7.5
        # Standardize stage coordinate system to 13.333 x 7.5
        self.slide_width = 13.333
        self.slide_height = 7.5
        self.scale_x = 13.333 / self.orig_width if self.orig_width > 0 else 1.0
        self.scale_y = 7.5 / self.orig_height if self.orig_height > 0 else 1.0

    @staticmethod
    def _iter_shapes(
        shapes, parent_left: float = 0.0, parent_top: float = 0.0
    ) -> Generator[Tuple[Any, float, float], None, None]:
        """Yield (shape, offset_left, offset_top) recursively descending groups.
        Enterprise decks heavily nest elements inside GroupShapes; a flat
        iteration over slide.shapes silently drops those children.
        """
        for shape in shapes:
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP or hasattr(shape, "shapes"):
                grp_left = shape.left.inches if shape.left else 0.0
                grp_top = shape.top.inches if shape.top else 0.0
                yield from SlideDecompiler._iter_shapes(
                    shape.shapes,
                    parent_left + grp_left,
                    parent_top + grp_top,
                )
            else:
                yield shape, parent_left, parent_top

    def _extract_slide_bg_color(self, slide) -> str:
        default_hex = "#FFFFFF" if self.raw_mode else "#001F33"
        try:
            # 1. Slide background
            if slide.background and slide.background.fill:
                fill = slide.background.fill
                if fill.type == MSO_FILL.SOLID and hasattr(fill, "fore_color"):
                    c = _color_to_hex(fill.fore_color)
                    if c:
                        return c
            # 2. Slide layout background
            if hasattr(slide, "slide_layout") and slide.slide_layout:
                layout = slide.slide_layout
                if layout.background and layout.background.fill:
                    fill = layout.background.fill
                    if fill.type == MSO_FILL.SOLID and hasattr(fill, "fore_color"):
                        c = _color_to_hex(fill.fore_color)
                        if c:
                            return c
            # 3. Slide master background
            if hasattr(slide, "slide_layout") and hasattr(slide.slide_layout, "slide_master"):
                master = slide.slide_layout.slide_master
                if master and master.background and master.background.fill:
                    fill = master.background.fill
                    if fill.type == MSO_FILL.SOLID and hasattr(fill, "fore_color"):
                        c = _color_to_hex(fill.fore_color)
                        if c:
                            return c
            # 4. Check for full-bleed background shape on slide
            for shp in slide.shapes:
                try:
                    s_left = shp.left.inches if shp.left is not None else 0.0
                    s_top = shp.top.inches if shp.top is not None else 0.0
                    s_w = shp.width.inches if shp.width is not None else 0.0
                    s_h = shp.height.inches if shp.height is not None else 0.0
                    if s_left <= 0.2 and s_top <= 0.2 and s_w >= (self.orig_width - 0.5) and s_h >= (self.orig_height - 0.5):
                        if shp.fill and shp.fill.type == MSO_FILL.SOLID and hasattr(shp.fill, "fore_color"):
                            c = _color_to_hex(shp.fill.fore_color)
                            if c:
                                return c
                except Exception:
                    pass
        except Exception:
            pass
        return default_hex

    @staticmethod
    def _resolve_shape_bounds(shape, slide=None) -> Tuple[float, float, float, float]:
        """Resolve (left_in, top_in, width_in, height_in) in inches, resolving inherited placeholder bounds."""
        raw_l = shape.left.inches if (hasattr(shape, "left") and shape.left is not None) else None
        raw_t = shape.top.inches if (hasattr(shape, "top") and shape.top is not None) else None
        raw_w = shape.width.inches if (hasattr(shape, "width") and shape.width is not None) else None
        raw_h = shape.height.inches if (hasattr(shape, "height") and shape.height is not None) else None

        is_missing = (
            raw_l is None or raw_t is None or raw_w is None or raw_h is None
            or (raw_l == 0.0 and raw_t == 0.0 and raw_w == 0.0 and raw_h == 0.0)
        )

        if is_missing and slide is not None and getattr(shape, "is_placeholder", False):
            try:
                ph_fmt = shape.placeholder_format
                p_idx = getattr(ph_fmt, "idx", None)
                p_type = getattr(ph_fmt, "type", None)
                layout = getattr(slide, "slide_layout", None)
                if layout and hasattr(layout, "placeholders"):
                    match_ph = None
                    for l_ph in layout.placeholders:
                        l_fmt = getattr(l_ph, "placeholder_format", None)
                        if l_fmt and l_fmt.idx == p_idx:
                            match_ph = l_ph
                            break
                    if match_ph is None and p_type is not None:
                        for l_ph in layout.placeholders:
                            l_fmt = getattr(l_ph, "placeholder_format", None)
                            if l_fmt and l_fmt.type == p_type:
                                match_ph = l_ph
                                break
                    if match_ph is not None:
                        if raw_l is None or raw_l == 0.0:
                            raw_l = match_ph.left.inches if match_ph.left is not None else 0.5
                        if raw_t is None or raw_t == 0.0:
                            raw_t = match_ph.top.inches if match_ph.top is not None else 0.5
                        if raw_w is None or raw_w == 0.0:
                            raw_w = match_ph.width.inches if match_ph.width is not None else 2.0
                        if raw_h is None or raw_h == 0.0:
                            raw_h = match_ph.height.inches if match_ph.height is not None else 1.0
            except Exception:
                pass

        left_in = raw_l if raw_l is not None else 0.5
        top_in = raw_t if raw_t is not None else 0.5
        width_in = raw_w if raw_w is not None else 2.0
        height_in = raw_h if raw_h is not None else 1.0
        return left_in, top_in, width_in, height_in

    def parse_deck(self) -> List[Dict[str, Any]]:
        parsed_slides = []
        for index, slide in enumerate(self.prs.slides):
            slide_text_accumulator = []
            nodes = []

            # Extract Slide Background Image if present
            bg_image_base64 = None
            try:
                if slide.background and slide.background.fill:
                    if slide.background.fill.type == 6:  # PICTURE
                        image_blob = slide.background.fill.picture.image.blob
                        bg_image_base64 = base64.b64encode(image_blob).decode('utf-8')
            except Exception:
                pass

            smartart_processed = False
            for shape, off_left, off_top in self._iter_shapes(slide.shapes):
                # SmartArt Graphic Frame detection
                is_gf = hasattr(shape, "_element") and getattr(shape._element, "tag", "").endswith("graphicFrame")
                if is_gf:
                    if not getattr(shape, "has_table", False) and not getattr(shape, "has_chart", False):
                        gf_l = round((shape.left.inches if shape.left else 0.8) * self.scale_x, 3)
                        gf_t = round((shape.top.inches if shape.top else 1.55) * self.scale_y, 3)
                        gf_w = round((shape.width.inches if shape.width else 11.73) * self.scale_x, 3)
                        gf_h = round((shape.height.inches if shape.height else 4.9) * self.scale_y, 3)
                        sa_nodes = self._extract_smartart_nodes(slide, gf_l, gf_t, gf_w, gf_h)
                        if sa_nodes:
                            smartart_processed = True
                            for sa_node in sa_nodes:
                                nodes.append(sa_node)
                                if sa_node.get("type") == "TEXT_CONTAINER" and sa_node.get("paragraphs"):
                                    for p in sa_node["paragraphs"]:
                                        if p.get("text"):
                                            slide_text_accumulator.append(p["text"])
                            continue

                node = self._extract_shape_data(shape, off_left, off_top, slide=slide)
                if node:
                    nodes.append(node)
                    is_ft = False
                    if _is_footer_or_slidenum and _is_footer_or_slidenum(shape):
                        is_ft = True
                    if not is_ft and node.get("type") == "TEXT_CONTAINER" and node.get("paragraphs"):
                        for p in node["paragraphs"]:
                            if p.get("text"):
                                slide_text_accumulator.append(p["text"])

            if not smartart_processed:
                try:
                    has_diagram = any("drawing" in str(getattr(r, "target_ref", "")).lower() or "diagram" in str(getattr(r, "reltype", "")).lower() for r in slide.part.rels.values())
                    if has_diagram:
                        sa_nodes = self._extract_smartart_nodes(slide, 0.80, 1.55, 11.73, 4.90)
                        if sa_nodes:
                            for sa_node in sa_nodes:
                                nodes.append(sa_node)
                                if sa_node.get("type") == "TEXT_CONTAINER" and sa_node.get("paragraphs"):
                                    for p in sa_node["paragraphs"]:
                                        if p.get("text"):
                                            slide_text_accumulator.append(p["text"])
                except Exception:
                    pass

            # Layer order: PICTURES at bottom, containers in middle, text on top
            def _layer_priority(n: Dict[str, Any]) -> int:
                t = n.get("type")
                if t in ("PICTURE", "IMAGE"):
                    return 0
                if t in ("RECTANGLE_SHAPE", "CONTAINER_SHAPE", "OVAL_SHAPE", "LINE_SHAPE", "ARROW_SHAPE"):
                    return 1
                if t in ("TABLE", "CHART"):
                    return 2
                return 3

            nodes.sort(key=_layer_priority)

            full_text = " ".join(slide_text_accumulator)

            archetype_name = "Original Deck" if self.raw_mode else "Multi-Card Content"
            if not self.raw_mode and SlideGraphClassifier:
                try:
                    archetype = SlideGraphClassifier.classify(slide, index, len(self.prs.slides))
                    archetype_name = archetype.value
                except Exception:
                    pass

            bg_color = self._extract_slide_bg_color(slide)

            slide_ast = {
                "slide_id": index + 1,
                "archetype": archetype_name,
                "dimensions": {"width": self.slide_width, "height": self.slide_height},
                "background_color": bg_color,
                "full_text": full_text,
                "bg_image": bg_image_base64,
                "nodes": nodes
            }
            parsed_slides.append(slide_ast)

        return parsed_slides

    def _extract_smartart_nodes(
        self, slide, gf_left: float = 0.80, gf_top: float = 1.55, gf_w: float = 11.73, gf_h: float = 4.90
    ) -> List[Dict[str, Any]]:
        """Parse SmartArt diagram XML and emit SHAPE and TEXT AST nodes."""
        nodes = []
        try:
            drawing_part = None
            data_part = None
            for r in slide.part.rels.values():
                t = str(getattr(r, "target_ref", "")).lower()
                rel_type = str(getattr(r, "reltype", "")).lower()
                if "diagramdrawing" in rel_type or "drawing" in t:
                    drawing_part = r.target_part
                if "diagramdata" in rel_type or "data" in t:
                    data_part = r.target_part

            node_id = 8000
            if drawing_part and hasattr(drawing_part, "blob"):
                from lxml import etree
                root = etree.fromstring(drawing_part.blob)
                sp_list = root.xpath('.//*[local-name()="sp"]')
                for sp in sp_list:
                    tx = ''.join(sp.xpath('.//*[local-name()="t"]/text()')).strip()
                    off = sp.xpath('.//*[local-name()="off"]')
                    ext = sp.xpath('.//*[local-name()="ext"]')
                    if not off or not ext:
                        continue
                    x_in = int(off[0].get("x", "0")) / 914400.0
                    y_in = int(off[0].get("y", "0")) / 914400.0
                    w_in = int(ext[0].get("cx", "0")) / 914400.0
                    h_in = int(ext[0].get("cy", "0")) / 914400.0

                    node_x = round(gf_left + x_in * self.scale_x, 3)
                    node_y = round(gf_top + y_in * self.scale_y, 3)
                    node_w = round(max(_MIN_EXTENT, w_in * self.scale_x), 3)
                    node_h = round(max(_MIN_EXTENT, h_in * self.scale_y), 3)

                    if tx:
                        node_id += 1
                        is_sub = len(tx) < 35
                        f_color = "#18181B" if self.raw_mode else ("#FFB81C" if is_sub else "#FFFFFF")
                        f_face = "Calibri" if self.raw_mode else "Arial"
                        f_size = 13.0 if is_sub else 11.0
                        nodes.append({
                            "id": node_id,
                            "type": "TEXT_CONTAINER",
                            "shape_name": "SmartArtText",
                            "x": node_x,
                            "y": node_y,
                            "w": node_w,
                            "h": node_h,
                            "bbox": {"left": node_x, "top": node_y, "width": node_w, "height": node_h},
                            "paragraphs": [{
                                "text": tx,
                                "level": 0,
                                "font_size": f_size,
                                "font_face": f_face,
                                "font_color": f_color,
                                "bold": is_sub if not self.raw_mode else False,
                                "italic": False,
                                "align": "LEFT",
                                "runs": [{"text": tx, "color": f_color, "font_name": f_face, "font_size": f_size, "bold": is_sub, "italic": False}],
                            }],
                            "is_title": False,
                        })
                    elif w_in > 0.05 and h_in > 0.05:
                        node_id += 1
                        nodes.append({
                            "id": node_id,
                            "type": "RECTANGLE_SHAPE",
                            "shape_name": "SmartArtShape",
                            "x": node_x,
                            "y": node_y,
                            "w": node_w,
                            "h": node_h,
                            "bbox": {"left": node_x, "top": node_y, "width": node_w, "height": node_h},
                            "fill_color": "#E2E8F0" if self.raw_mode else "#081E36",
                            "line_color": "#CBD5E1" if self.raw_mode else "#1E3A5F",
                        })

            # Fallback if drawing_part had no nodes: extract from data_part
            if not nodes and data_part and hasattr(data_part, "blob"):
                from lxml import etree
                root = etree.fromstring(data_part.blob)
                texts = [t.text.strip() for t in root.xpath('.//*[local-name()="t"]') if t.text and t.text.strip()]
                if texts:
                    step_w = gf_w / len(texts)
                    for i, txt in enumerate(texts):
                        node_id += 1
                        step_x = round(gf_left + i * step_w, 3)
                        f_color = "#18181B" if self.raw_mode else "#FFB81C"
                        nodes.append({
                            "id": node_id,
                            "type": "TEXT_CONTAINER",
                            "shape_name": "SmartArtStep",
                            "x": step_x,
                            "y": gf_top,
                            "w": round(step_w * 0.9, 3),
                            "h": gf_h,
                            "bbox": {"left": step_x, "top": gf_top, "width": round(step_w * 0.9, 3), "height": gf_h},
                            "paragraphs": [{
                                "text": txt,
                                "level": 0,
                                "font_size": 12.0,
                                "font_face": "Calibri" if self.raw_mode else "Arial",
                                "font_color": f_color,
                                "bold": True,
                                "italic": False,
                                "align": "CENTER",
                                "runs": [{"text": txt, "color": f_color, "font_name": "Arial", "font_size": 12.0, "bold": True, "italic": False}],
                            }],
                            "is_title": False,
                        })
        except Exception:
            pass
        return nodes

    def _extract_shape_data(
        self, shape, offset_left: float = 0.0, offset_top: float = 0.0, slide=None
    ) -> Dict[str, Any] | None:
        try:
            try:
                raw_l, raw_t, raw_w, raw_h = self._resolve_shape_bounds(shape, slide)
                raw_left = raw_l + offset_left
                raw_top = raw_t + offset_top
                raw_width = raw_w
                raw_height = raw_h

                left = max(0.0, round(raw_left * self.scale_x, 3))
                top = max(0.0, round(raw_top * self.scale_y, 3))
                width = max(_MIN_EXTENT, round(raw_width * self.scale_x, 3))
                height = max(_MIN_EXTENT, round(raw_height * self.scale_y, 3))
            except Exception:
                left, top, width, height = 0.5, 0.5, 2.0, 1.0

            shape_id = getattr(shape, "shape_id", 0)
            shape_type = getattr(shape, "shape_type", None)

            shape_name_str = getattr(shape, "name", "") or ""
            base_data: Dict[str, Any] = {
                "id": shape_id,

                "name": shape_name_str,
                "x": left,
                "y": top,
                "w": width,
                "h": height,
                "bbox": {
                    "left": left,
                    "top": top,
                    "width": width,
                    "height": height
                }
            }

            # Fill and Line properties
            fill_hex = None
            try:
                if shape.fill and shape.fill.type == MSO_FILL.SOLID and hasattr(shape.fill, "fore_color"):
                    fill_hex = _color_to_hex(shape.fill.fore_color)
            except Exception:
                pass
            if fill_hex:
                base_data["fill_hex"] = fill_hex
                base_data["fill_color"] = fill_hex

            line_hex = None
            try:
                if shape.line and hasattr(shape.line, "color") and hasattr(shape.line.color, "rgb"):
                    line_hex = _color_to_hex(shape.line.color)
            except Exception:
                pass
            if line_hex:
                base_data["line_hex"] = line_hex
                base_data["line_color"] = line_hex

            # Check for bottom frame
            if "BrandMorphBottomFrame" in shape_name_str or "BrandMorphAccentLine" in shape_name_str:
                base_data["type"] = "RECTANGLE_SHAPE"
                base_data["shape_name"] = "BrandMorphBottomFrame"
                return base_data

            # 1. Check for Picture / Image FIRST (so picture placeholders are not swallowed by has_text_frame)
            is_picture = False
            image_bytes = None
            content_type = "image/png"
            try:
                if shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                    if getattr(shape, "image", None) is not None:
                        is_picture = True
                        image_bytes = shape.image.blob
                        content_type = getattr(shape.image, "content_type", None) or f"image/{getattr(shape.image, 'ext', 'png')}"
            except Exception:
                pass

            if not is_picture:
                try:
                    # Check for picture placeholder or blip in shape element
                    blips = shape.element.xpath('.//a:blip')
                    if blips:
                        r_id = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if r_id and hasattr(shape, "part") and r_id in shape.part.related_parts:
                            rel = shape.part.related_parts[r_id]
                            if hasattr(rel, "blob"):
                                image_bytes = rel.blob
                                content_type = getattr(rel, "content_type", "image/png")
                                is_picture = True
                except Exception:
                    pass

            if is_picture and image_bytes:
                try:
                    image_base64 = base64.b64encode(image_bytes).decode('utf-8')
                    data_uri = f"data:{content_type};base64,{image_base64}"
                    base_data["type"] = "PICTURE"
                    base_data["image_data_uri"] = data_uri
                    base_data["image_src"] = data_uri
                    base_data["image_data"] = data_uri
                    return base_data
                except Exception:
                    pass

            # 2. Text Container Frame
            if getattr(shape, "has_text_frame", False) and shape.text_frame and shape.text_frame.text.strip():
                is_footer = False
                if _is_footer_or_slidenum and _is_footer_or_slidenum(shape):
                    is_footer = True

                paragraphs = []
                is_title_frame = False if is_footer else (
                    (getattr(shape, "is_placeholder", False) and getattr(shape, "placeholder_format", None) and shape.placeholder_format.type in (1, 3))
                    or ("title" in (getattr(shape, "name", "") or "").lower() and "subtitle" not in (getattr(shape, "name", "") or "").lower())
                    or (top < 1.8 and height < 2.0 and not is_footer)
                )

                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if not self.raw_mode and ("Font Guidance" in txt or "Do not change the template" in txt):
                        continue

                    runs_data = []
                    first_run_color = None
                    first_run_font = None
                    first_run_size = None
                    first_run_bold = None
                    first_run_italic = None

                    for r in p.runs:
                        r_color = None
                        try:
                            if r.font and r.font.color and r.font.color.type is not None:
                                r_color = _color_to_hex(r.font.color)
                        except Exception:
                            pass
                        if r_color and not first_run_color:
                            first_run_color = r_color
                        if r.font and r.font.name and not first_run_font:
                            first_run_font = r.font.name
                        if r.font and r.font.size and r.font.size.pt and not first_run_size:
                            first_run_size = r.font.size.pt
                        if r.font and r.font.bold is not None and first_run_bold is None:
                            first_run_bold = bool(r.font.bold)
                        if r.font and r.font.italic is not None and first_run_italic is None:
                            first_run_italic = bool(r.font.italic)

                        runs_data.append({
                            "text": r.text,
                            "font_name": r.font.name if r.font else None,
                            "font_size": r.font.size.pt if r.font and r.font.size else None,
                            "bold": bool(r.font.bold) if r.font else False,
                            "italic": bool(r.font.italic) if r.font else False,
                            "color": r_color,
                        })

                    p_align = 'LEFT'
                    try:
                        if str(p.alignment) == 'CENTER (2)':
                            p_align = 'CENTER'
                        elif str(p.alignment) == 'RIGHT (3)':
                            p_align = 'RIGHT'
                        elif str(p.alignment) == 'JUSTIFY (4)':
                            p_align = 'JUSTIFY'
                    except Exception:
                        pass

                    if is_footer:
                        if self.raw_mode:
                            font_color = first_run_color or "#64748B"
                            font_name = first_run_font or "Calibri"
                            font_size = first_run_size or 8.5
                            bold = False
                            italic = False
                        else:
                            font_color = "#94A3B8"
                            font_name = "Arial"
                            font_size = 8.5
                            bold = False
                            italic = False
                    elif self.raw_mode:
                        bg_color = self._extract_slide_bg_color(slide) if slide else "#FFFFFF"
                        bg_lum = 1.0
                        try:
                            clean_bg = bg_color.lstrip("#")
                            if len(clean_bg) == 6:
                                r, g, b = int(clean_bg[0:2], 16)/255.0, int(clean_bg[2:4], 16)/255.0, int(clean_bg[4:6], 16)/255.0
                                bg_lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
                        except Exception:
                            bg_lum = 1.0
                        default_text_color = "#FFFFFF" if bg_lum < 0.35 else "#18181B"

                        p_color = _color_to_hex(p.font.color) if (p.font and hasattr(p.font, "color")) else None
                        font_color = first_run_color or p_color or default_text_color
                        font_name = first_run_font or (p.font.name if p.font else None) or "Calibri"
                        font_size = first_run_size or (p.font.size.pt if p.font and p.font.size else (28.0 if is_title_frame else 14.0))
                        bold = bool(p.font.bold) if (p.font and p.font.bold is not None) else (first_run_bold or False)
                        italic = bool(p.font.italic) if (p.font and p.font.italic is not None) else (first_run_italic or False)
                    else:
                        is_sub = (first_run_bold or (first_run_size and 13.0 <= first_run_size <= 16.0) or (len(shape.text_frame.paragraphs) > 1 and p == shape.text_frame.paragraphs[0] and len(txt) <= 60))
                        if is_title_frame:
                            font_color = "#FFB81C"
                            font_name = "Times New Roman"
                            font_size = p.font.size.pt if p.font and p.font.size else 28.0
                            bold = False
                        elif is_sub:
                            font_color = first_run_color if first_run_color in ("#FFB81C", "#00FFFF", "#00E5FF") else "#FFB81C"
                            font_name = "Arial"
                            font_size = p.font.size.pt if p.font and p.font.size else 14.0
                            bold = True
                        else:
                            font_color = first_run_color or "#FFFFFF"
                            font_name = "Arial"
                            font_size = p.font.size.pt if p.font and p.font.size else 11.0
                            bold = False
                        italic = bool(p.font.italic) if p.font else False


                    paragraphs.append({
                        "text": p.text,
                        "level": getattr(p, "level", 0),
                        "font_size": font_size,
                        "font_face": font_name,
                        "font_color": font_color,
                        "bold": bold,
                        "italic": italic,
                        "underline": bool(p.font.underline) if p.font else False,
                        "align": p_align,
                        "runs": runs_data
                    })

                base_data["type"] = "TEXT_CONTAINER"
                base_data["paragraphs"] = paragraphs
                base_data["is_title"] = is_title_frame
                return base_data

            # 3. Geometric Shapes / Arrows / Containers
            if shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
                shape_enum_str = str(getattr(shape, 'auto_shape_type', 'RECTANGLE'))
                if "ARROW" in shape_enum_str or "CHEVRON" in shape_enum_str:
                    base_data["type"] = "ARROW_SHAPE"
                elif "OVAL" in shape_enum_str or "CIRCLE" in shape_enum_str:
                    base_data["type"] = "OVAL_SHAPE"
                else:
                    base_data["type"] = "RECTANGLE_SHAPE"

                base_data["shape_name"] = shape_enum_str
                return base_data

            # 4. Lines and Connectors
            tag = getattr(shape, "_element", None)
            is_connector = tag is not None and getattr(tag, "tag", "").endswith("}cxnSp")
            if shape_type == MSO_SHAPE_TYPE.LINE or is_connector:
                base_data["type"] = "LINE_SHAPE"
                base_data["shape_name"] = "CONNECTOR_LINE" if is_connector else "LINE"
                return base_data

            # 5. Tables
            is_table = shape_type == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False)
            if is_table:
                base_data["type"] = "TABLE"
                base_data["shape_name"] = "TABLE_SHAPE"
                try:
                    table = shape.table
                    rows_data = []
                    total_w = sum(col.width.inches for col in table.columns if col.width) if table.columns else 0.0
                    if total_w > 0:
                        col_widths = [round((col.width.inches / total_w) * 100, 1) for col in table.columns]
                    else:
                        n_cols = len(table.columns) if table.columns else 1
                        col_widths = [round(100.0 / n_cols, 1)] * n_cols

                    for r_idx, row in enumerate(table.rows):
                        row_cells = []
                        for c_idx, cell in enumerate(row.cells):
                            text = cell.text.strip() if cell.text else ""

                            # Extract cell fill
                            fill_color = None
                            try:
                                if cell.fill and cell.fill.type == 1 and hasattr(cell.fill, "fore_color"):
                                    fill_color = f"#{str(cell.fill.fore_color.rgb)}".upper()
                            except Exception:
                                pass
                            if not fill_color:
                                if self.raw_mode:
                                    fill_color = "#FFFFFF"
                                else:
                                    fill_color = "#FFB81C" if r_idx == 0 else ("#081E36" if r_idx % 2 == 1 else "#030C1E")

                            # Extract cell text styling
                            text_color = None
                            bold = (r_idx == 0 or c_idx == 0)
                            font_size = 12.0 if r_idx == 0 else 11.0
                            align = "left"

                            if cell.text_frame and cell.text_frame.paragraphs:
                                p0 = cell.text_frame.paragraphs[0]
                                if p0.alignment == PP_ALIGN.CENTER:
                                    align = "center"
                                elif p0.alignment == PP_ALIGN.RIGHT:
                                    align = "right"
                                else:
                                    align = "left"

                                for p in cell.text_frame.paragraphs:
                                    for r in p.runs:
                                        if r.font:
                                            if r.font.bold is not None:
                                                bold = r.font.bold
                                            if r.font.size and r.font.size.pt:
                                                font_size = r.font.size.pt
                                            if r.font.color and r.font.color.type == 1:
                                                try:
                                                    text_color = f"#{str(r.font.color.rgb)}".upper()
                                                except Exception:
                                                    pass
                                        if text_color:
                                            break
                                    if text_color:
                                        break

                            if not text_color:
                                if self.raw_mode:
                                    bg_color = self._extract_slide_bg_color(slide) if slide else "#FFFFFF"
                                    bg_lum = 1.0
                                    try:
                                        clean_bg = bg_color.lstrip("#")
                                        if len(clean_bg) == 6:
                                            r, g, b = int(clean_bg[0:2], 16)/255.0, int(clean_bg[2:4], 16)/255.0, int(clean_bg[4:6], 16)/255.0
                                            bg_lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
                                    except Exception:
                                        bg_lum = 1.0
                                    text_color = "#FFFFFF" if bg_lum < 0.35 else "#18181B"
                                else:
                                    text_color = "#001F33" if r_idx == 0 else ("#FFB81C" if c_idx == 0 else "#FFFFFF")


                            row_cells.append({
                                "text": text,
                                "fill_color": fill_color,
                                "text_color": text_color,
                                "bold": bold,
                                "font_size": font_size,
                                "align": align,
                            })
                        rows_data.append(row_cells)

                    base_data["table_data"] = {
                        "rows": rows_data,
                        "col_widths": col_widths,
                        "num_rows": len(table.rows),
                        "num_cols": len(table.columns),
                    }
                except Exception:
                    pass
                return base_data

            # 6. Charts
            is_chart = shape_type == MSO_SHAPE_TYPE.CHART or getattr(shape, "has_chart", False)
            if is_chart:
                base_data["type"] = "CHART"
                base_data["shape_name"] = "CHART_SHAPE"
                return base_data

            return None
        except Exception:
            return None
