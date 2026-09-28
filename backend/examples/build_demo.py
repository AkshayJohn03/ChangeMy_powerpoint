"""Build a committed before/after demo for BrandMorph v2.

Run from backend/:  python examples/build_demo.py

Creates examples/demo_raw.pptx (a deterministically "off-brand" 3-slide deck:
generic corporate blue, mismatched fonts, hand-painted colors, a chart, a
table and a grouped shape), then applies the Akkodis brand in place and
writes examples/demo_akkodis_branded.pptx + examples/report/morph_report.md.
The branded output is committed so reviewers can open both decks side by side.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from brandmorph.brand.profile import BrandDNA  # noqa: E402
from brandmorph.engine.pipeline import MorphOptions, MorphPipeline  # noqa: E402
from brandmorph.report import MorphReport  # noqa: E402
from brandmorph.engine.analyzer import analyze  # noqa: E402

EXAMPLES = Path(__file__).resolve().parent
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

OLD_BG = RGBColor(0x1F, 0x4E, 0x79)      # generic corporate blue
OLD_CARD = RGBColor(0x2E, 0x75, 0xB6)    # lighter blue card
OLD_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
OLD_ACCENT = RGBColor(0xC0, 0x00, 0x00)  # hand-painted red
OLD_MUTED = RGBColor(0xBF, 0xBF, 0xBF)


def _run(p, text, size, color, name, bold=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = name
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def _box(slide, x, y, w, h, text, size, color, name, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    _run(tb.text_frame.paragraphs[0], text, size, color, name, bold)
    return tb


def build_demo(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    # ---- slide 1: cover -------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = OLD_BG; bg.line.fill.background()
    _box(s, Inches(0.9), Inches(2.2), Inches(10), Inches(1.1),
         "FY25 Operations Review", 40, OLD_TEXT, "Times New Roman", bold=True)
    _box(s, Inches(0.9), Inches(3.4), Inches(10), Inches(0.6),
         "Prepared by the Delivery Excellence Office", 18, OLD_MUTED, "Georgia")

    # ---- slide 2: KPIs + table -------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = OLD_BG; bg.line.fill.background()
    _box(s, Inches(0.8), Inches(0.5), Inches(9), Inches(0.8),
         "Service Performance Summary", 30, OLD_TEXT, "Verdana", bold=True)
    for i, (label, value, color) in enumerate([
        ("On-time delivery", "97.6%", OLD_TEXT),
        ("First-time right", "96.8%", OLD_TEXT),
        ("Margin alert", "-2.1%", OLD_ACCENT),
    ]):
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(0.8 + i * 3.4), Inches(1.7),
                                  Inches(3.0), Inches(1.6))
        card.fill.solid(); card.fill.fore_color.rgb = OLD_CARD; card.line.fill.background()
        tf = card.text_frame
        _run(tf.paragraphs[0], value, 28, color, "Verdana", bold=True)
        _run(tf.add_paragraph(), label, 13, OLD_MUTED, "Verdana")
    table = s.shapes.add_table(3, 3, Inches(0.8), Inches(3.8), Inches(9.4), Inches(1.8)).table
    for c, head in enumerate(["Region", "OTD", "Churn"]):
        cell = table.cell(0, c)
        cell.text = head
        cell.fill.solid(); cell.fill.fore_color.rgb = OLD_CARD
        for p in cell.text_frame.paragraphs:
            for r in p.runs or [p.add_run()]:
                r.font.color.rgb = OLD_TEXT
    for r_i, row in enumerate([["AMER", "96.1%", "1.9%"], ["EMEA", "98.4%", "1.2%"]], start=1):
        for c, val in enumerate(row):
            table.cell(r_i, c).text = val
    _box(s, Inches(0.8), Inches(6.1), Inches(9), Inches(0.4),
         "Source: internal ops dashboard, FY25 Q3", 12, OLD_MUTED, "Arial")
    group = s.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(11.0), Inches(6.4), Inches(1.8), Inches(0.4))
    _run(inner.text_frame.paragraphs[0], "Confidential", 11, OLD_MUTED, "Arial")

    # ---- slide 3: chart ---------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = OLD_BG; bg.line.fill.background()
    _box(s, Inches(0.8), Inches(0.5), Inches(9), Inches(0.8),
         "Quarterly Trend", 30, OLD_TEXT, "Verdana", bold=True)
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3", "Q4"]
    cd.add_series("OTD", (94.1, 96.0, 97.2, 97.6))
    cd.add_series("FTR", (92.8, 94.5, 96.1, 96.8))
    gframe = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,
                                Inches(0.8), Inches(1.6), Inches(8.5), Inches(5.0), cd)
    for series in gframe.chart.series:
        line = series.format.line
        line.color.rgb = OLD_ACCENT if series is gframe.chart.series[0] else OLD_CARD
    _box(s, Inches(9.8), Inches(2.2), Inches(3), Inches(2.4),
         "Both series improved for four consecutive quarters; the Q4 margin "
         "dip needs a dedicated recovery plan with named owners.", 14, OLD_TEXT, "Arial")
    prs.save(str(path))


def main() -> int:
    raw = EXAMPLES / "demo_raw.pptx"
    branded = EXAMPLES / "demo_akkodis_branded.pptx"
    report_dir = EXAMPLES / "report"
    if not raw.exists():
        build_demo(raw)
        print(f"built {raw.name}")
    dna = BrandDNA.load("akkodis")
    pipeline = MorphPipeline(dna, options=MorphOptions())
    report: MorphReport = pipeline.run(raw, branded, report_dir=str(report_dir))
    before, after = analyze(raw), analyze(branded)
    fidelity = all(
        b[k] == a[k] for b, a in zip(before["slides"], after["slides"])
        for k in ("shape_count", "charts", "tables", "pictures", "text_shapes")
    )
    print(f"changes: {report.total_changes} "
          f"({report.counts_by_stage})")
    print(f"fidelity preserved: {fidelity}")
    print(f"review flags: {len(report.review_flags)}")
    print(f"wrote {branded.name} + report/morph_report.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
