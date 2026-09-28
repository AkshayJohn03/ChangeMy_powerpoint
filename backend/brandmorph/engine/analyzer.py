"""DeckAnalyzer: inventory the deck before touching it.

Every downstream stage consults this inventory, and the API exposes it so a
user can see "what will change" before committing (v1 gave no such signal —
content vanished silently, diagnosis defect D4).
"""

from __future__ import annotations

import hashlib
from collections import Counter
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as _PresentationCls

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _iter_shapes(shapes):
    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield shp
            yield from _iter_shapes(shp.shapes)
        else:
            yield shp


def analyze(path_or_prs) -> dict[str, Any]:
    """Build a JSON-serializable inventory of the deck."""
    prs = (
        path_or_prs
        if isinstance(path_or_prs, _PresentationCls)
        else Presentation(str(path_or_prs))
    )
    slides_out = []
    total_images: dict[str, str] = {}

    for idx, slide in enumerate(prs.slides):
        counts = Counter()
        explicit_colors: Counter = Counter()
        theme_refs: Counter = Counter()
        fonts: Counter = Counter()
        for shp in _iter_shapes(slide.shapes):
            counts["shapes"] += 1
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
                try:
                    blob = shp.image.blob
                    total_images[hashlib.sha256(blob).hexdigest()[:12]] = shp.image.sha1[:12]
                except Exception:
                    pass
            if getattr(shp, "has_chart", False) and shp.has_chart:
                counts["charts"] += 1
            if getattr(shp, "has_table", False) and shp.has_table:
                counts["tables"] += 1
            if shp.has_text_frame:
                counts["text_shapes"] += 1
        # XML-level color/font census (covers fills, lines, runs, tables)
        for el in slide._element.iter():
            tag = el.tag
            if tag == f"{{{_A}}}srgbClr" and el.get("val"):
                explicit_colors[el.get("val").upper()] += 1
            elif tag == f"{{{_A}}}schemeClr" and el.get("val"):
                theme_refs[el.get("val")] += 1
            elif tag == f"{{{_A}}}latin" and el.get("typeface"):
                fonts[el.get("typeface")] += 1
        slides_out.append({
            "index": idx,
            "shape_count": counts["shapes"],
            "pictures": counts["pictures"],
            "charts": counts["charts"],
            "tables": counts["tables"],
            "text_shapes": counts["text_shapes"],
            "explicit_colors": dict(explicit_colors.most_common(20)),
            "theme_refs": dict(theme_refs.most_common(20)),
            "fonts": dict(fonts.most_common(20)),
        })

    return {
        "slide_count": len(prs.slides),
        "slide_size_emu": [prs.slide_width, prs.slide_height],
        "slides": slides_out,
        "image_hashes": sorted(total_images.keys()),
    }
