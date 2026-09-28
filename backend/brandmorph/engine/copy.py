"""CopyToneRewriter: brand voice under hard length budgets.

Two deterministic passes always run (offline, no LLM needed):
  - glossary enforcement (preferred spellings/terminology, case-insensitive)
  - banned-word detection (flagged for review — never silently deleted)

The optional LLM pass rewrites copy to the brand voice ONLY under the
template budget rule: replacement length <= original length x 1.1 (design
§2.1 ⑥). Numbers, dates and capitalised proper nouns are frozen; if the LLM
violates the budget or drops frozen tokens, the original text is kept and the
skip is logged. Without an LLMClient this stage is a documented no-op — the
engine remains fully functional.
"""

from __future__ import annotations

import asyncio
import re
from typing import Protocol

from pptx import Presentation

from ..brand.profile import BrandDNA
from ..report import ChangeEntry, MorphReport

_BUDGET_FACTOR = 1.1
_MIN_LL_LEN = 40  # don't waste LLM calls on labels
_NUMBER = re.compile(r"\d")


class LLMClient(Protocol):
    async def complete(self, prompt: str) -> str: ...


def _replace_paragraph_text(p, new_text: str) -> None:
    """First-run replace: keeps run formatting, drops the other runs."""
    runs = p.runs
    if not runs:
        p.add_run().text = new_text
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


class CopyToneRewriter:
    def __init__(self, dna: BrandDNA, llm: LLMClient | None = None):
        self.dna = dna
        self.llm = llm

    def apply(self, prs: Presentation, report: MorphReport) -> None:
        self._glossary_pass(prs, report)
        self._banned_pass(prs, report)
        if self.llm is not None:
            asyncio.run(self._llm_pass(prs, report))

    # -- deterministic passes ---------------------------------------------
    def _glossary_pass(self, prs: Presentation, report: MorphReport) -> None:
        glossary = self.dna.tone.glossary
        if not glossary:
            return
        for idx, slide in enumerate(prs.slides):
            for shp in slide.shapes:
                if not (getattr(shp, "has_text_frame", False) and shp.has_text_frame):
                    continue
                for p in shp.text_frame.paragraphs:
                    original = p.text
                    if not original.strip():
                        continue
                    new = original
                    for wrong, right in glossary.items():
                        pattern = r"\b" + re.escape(wrong) + r"\b"
                        new = re.sub(pattern, right, new, flags=re.IGNORECASE)
                    if new != original:
                        _replace_paragraph_text(p, new)
                        report.add(ChangeEntry(
                            stage="copy", slide=idx,
                            element=shp.name, kind="text",
                            before=original, after=new, role="glossary",
                        ))

    def _banned_pass(self, prs: Presentation, report: MorphReport) -> None:
        banned = [b.lower() for b in self.dna.tone.banned_words]
        if not banned:
            return
        for idx, slide in enumerate(prs.slides):
            for shp in slide.shapes:
                if not (getattr(shp, "has_text_frame", False) and shp.has_text_frame):
                    continue
                text = shp.text_frame.text.lower()
                for word in banned:
                    if re.search(r"\b" + re.escape(word) + r"\b", text):
                        report.flag(
                            f"slide {idx}: banned word '{word}' found in '{shp.name}' — "
                            f"manual rewrite required (not auto-deleted)"
                        )

    # -- optional LLM pass --------------------------------------------------
    async def _llm_pass(self, prs: Presentation, report: MorphReport) -> None:
        voice = self.dna.tone.voice or "professional, concise"
        for idx, slide in enumerate(prs.slides):
            for shp in slide.shapes:
                if not (getattr(shp, "has_text_frame", False) and shp.has_text_frame):
                    continue
                for p in shp.text_frame.paragraphs:
                    original = p.text
                    if len(original.strip()) < _MIN_LL_LEN:
                        continue
                    budget = int(len(original) * _BUDGET_FACTOR)
                    frozen = re.findall(r"[\w][\w'.&-]*", original)
                    frozen = [t for t in frozen if _NUMBER.search(t) or t[:1].isupper()]
                    prompt = (
                        f"Rewrite this presentation copy in the brand voice "
                        f"({voice}). HARD RULES: max {budget} characters; keep every "
                        f"number, date and proper noun exactly as-is "
                        f"({', '.join(frozen) if frozen else 'none'}); no quotes, "
                        f"no commentary, output only the rewritten text.\n\n{original}"
                    )
                    try:
                        rewritten = (await self.llm.complete(prompt)).strip().strip('"')
                    except Exception as exc:  # LLM failure must never corrupt a deck
                        report.flag(f"slide {idx}: LLM rewrite skipped ({exc})")
                        continue
                    if not rewritten or len(rewritten) > budget:
                        report.add(ChangeEntry(
                            stage="copy", slide=idx, element=shp.name, kind="text",
                            before=original, after="(kept original)",
                            note=f"LLM over budget ({len(rewritten)}>{budget})",
                        ))
                        continue
                    missing = [t for t in frozen if t not in rewritten]
                    if missing:
                        report.add(ChangeEntry(
                            stage="copy", slide=idx, element=shp.name, kind="text",
                            before=original, after="(kept original)",
                            note=f"LLM dropped frozen tokens: {missing}",
                        ))
                        continue
                    _replace_paragraph_text(p, rewritten)
                    report.add(ChangeEntry(
                        stage="copy", slide=idx, element=shp.name, kind="text",
                        before=original, after=rewritten, role="tone",
                        confidence=0.8,
                    ))
