"""Shared fixtures: synthetic decks with planted styling (design §2.3).

Decks are built in-test with python-pptx so the suite is fully offline and
deterministic. The 'mixed' deck carries fidelity sentinels (chart, table,
group, picture) that v1's rebuild pipeline destroyed.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BACKEND = Path(__file__).resolve().parents[1]
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from brandmorph.brand.profile import BrandDNA

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


@pytest.fixture(scope="session")
def akkodis() -> BrandDNA:
    return BrandDNA.load("akkodis")


@pytest.fixture(scope="session")
def corporate_light() -> BrandDNA:
    return BrandDNA.load("corporate_light")


def _new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def _add_run(paragraph, text: str, size: int, color, name: str, bold: bool = False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = name
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def _scheme_fill(shape, scheme_val: str) -> None:
    """Replace a shape's explicit fill with a theme (schemeClr) reference."""
    spPr = shape._element.spPr
    solid = spPr.find(qn("a:solidFill"))
    if solid is None:
        solid = etree.SubElement(spPr, qn("a:solidFill"))
    for child in list(solid):
        solid.remove(child)
    sc = etree.SubElement(solid, qn("a:schemeClr"))
    sc.set("val", scheme_val)


def _scheme_run_color(run, scheme_val: str) -> None:
    rPr = run._r.get_or_add_rPr()
    solid = rPr.find(qn("a:solidFill"))
    if solid is None:
        solid = etree.SubElement(rPr, qn("a:solidFill"))
    for child in list(solid):
        solid.remove(child)
    sc = etree.SubElement(solid, qn("a:schemeClr"))
    sc.set("val", scheme_val)


def build_hardcoded(path: Path) -> Path:
    """Explicit colors/fonts everywhere — the case v1 handled worst (D2/D3)."""
    prs = _new_deck()
    s = prs.slides.add_slide(prs.slide_layouts[6])

    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x01, 0x02, 0x03)  # near-black, not brand
    bg.line.fill.background()
    bg.shadow.inherit = False

    title = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(9), Inches(1.0))
    _add_run(
        title.text_frame.paragraphs[0], "Quarterly Business Review", 28,
        RGBColor(0xFF, 0xFF, 0xFF), "Times New Roman", bold=True,
    )
    body = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8), Inches(1.2))
    _add_run(
        body.text_frame.paragraphs[0],
        "Revenue grew 42 percent across all regions and the team Utilize new tooling.",
        14, RGBColor(0xCC, 0xCC, 0xCC), "Arial",
    )
    accent = s.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(4), Inches(0.5))
    _add_run(
        accent.text_frame.paragraphs[0], "Margin alert", 20,
        RGBColor(0xFF, 0x00, 0x00), "Arial", bold=True,
    )
    onbrand = s.shapes.add_textbox(Inches(0.8), Inches(4.4), Inches(4), Inches(0.5))
    _add_run(
        onbrand.text_frame.paragraphs[0], "Already on brand", 16,
        RGBColor(0xFF, 0xB8, 0x1C), "Arial",
    )
    overflow = s.shapes.add_textbox(Inches(7.0), Inches(3.4), Inches(3.0), Inches(0.6))
    _add_run(
        overflow.text_frame.paragraphs[0],
        "Comprehensive analysis of operational efficiency improvements across every "
        "business unit and regional market segment during the entire fiscal year",
        24, RGBColor(0xFF, 0xFF, 0xFF), "Arial",
    )
    dim = s.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(6), Inches(0.5))
    _add_run(
        dim.text_frame.paragraphs[0], "Footer note, low emphasis", 11,
        RGBColor(0x55, 0x55, 0x55), "Arial",
    )
    prs.save(str(path))
    return path


def build_themed(path: Path) -> Path:
    """Theme-referenced deck: schemeClr fills/runs that the theme pass owns."""
    prs = _new_deck()
    s = prs.slides.add_slide(prs.slide_layouts[6])

    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    card.fill.solid()
    _scheme_fill(card, "accent1")
    card.text_frame.text = ""

    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(4), Inches(1))
    band.fill.solid()
    _scheme_fill(band, "dk2")

    tb = s.shapes.add_textbox(Inches(6), Inches(1), Inches(5), Inches(1))
    run = _add_run(tb.text_frame.paragraphs[0], "Theme-driven headline", 24,
                   RGBColor(0, 0, 0), "Consolas", bold=True)
    _scheme_run_color(run, "dk2")
    prs.save(str(path))
    return path


def build_mixed(path: Path) -> Path:
    """Fidelity sentinels: chart + table + group + picture (v1 destroyed all)."""
    prs = _new_deck()
    s = prs.slides.add_slide(prs.slide_layouts[6])

    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Revenue", (1.0, 2.0, 3.0))
    cd.add_series("Cost", (2.0, 1.5, 2.5))
    gframe = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(0.5), Inches(5), Inches(3), cd
    )
    pt = gframe.chart.series[0].format.fill
    pt.solid()
    pt.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)

    table = s.shapes.add_table(3, 3, Inches(6), Inches(0.5), Inches(6), Inches(2)).table
    cell = table.cell(0, 0)
    cell.text = "Region"
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x22, 0x33, 0x44)

    group = s.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(0.5))
    _add_run(inner.text_frame.paragraphs[0], "Grouped note", 12,
             RGBColor(0xDD, 0xDD, 0xDD), "Arial")

    img = io.BytesIO()
    try:
        from PIL import Image

        Image.new("RGB", (40, 40), (200, 60, 60)).save(img, format="PNG")
        img.seek(0)
        s.shapes.add_picture(img, Inches(9), Inches(4), Inches(1), Inches(1))
    except ImportError:
        pass  # picture sentinel skipped when Pillow absent

    prs.save(str(path))
    return path
