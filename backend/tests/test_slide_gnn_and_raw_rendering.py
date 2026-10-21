"""Verification tests for SlideGNN ONNX model and true Raw layout/crop rendering."""

import os
import time
from pathlib import Path
import numpy as np
import pytest
import pptx
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

import onnxruntime as ort
from backend.decompiler import SlideDecompiler
from backend.brandmorph.engine.archetype_engine import SlideGraphClassifier, SlideArchetype


def test_slide_gnn_models_exist():
    onnx_path = Path("backend/models/slide_gnn.onnx")
    safetensors_path = Path("backend/models/slide_gnn.safetensors")

    assert onnx_path.exists(), f"Missing ONNX model: {onnx_path}"
    assert safetensors_path.exists(), f"Missing safetensors model: {safetensors_path}"
    assert onnx_path.stat().st_size > 10000, "ONNX model file is too small"
    assert safetensors_path.stat().st_size > 10000, "Safetensors model file is too small"


def test_slide_gnn_onnx_inference_latency():
    onnx_path = Path("backend/models/slide_gnn.onnx")
    opts = ort.SessionOptions()
    opts.intra_op_num_threads = 1
    session = ort.InferenceSession(str(onnx_path), sess_options=opts, providers=["CPUExecutionProvider"])

    for n_nodes in [2, 5, 8, 14]:
        X = np.random.randn(n_nodes, 16).astype(np.float32)
        A = np.eye(n_nodes, dtype=np.float32)

        # Warmup
        session.run(None, {"node_features": X, "adjacency": A})

        # Measure 20 runs
        t0 = time.perf_counter()
        for _ in range(20):
            node_logits, slide_logits = session.run(None, {"node_features": X, "adjacency": A})
        avg_latency_ms = (time.perf_counter() - t0) / 20 * 1000

        assert node_logits.shape == (n_nodes, 9)
        assert slide_logits.shape == (1, 9)
        assert avg_latency_ms < 5.0, f"Latency {avg_latency_ms:.3f} ms exceeds 5.0 ms target for N={n_nodes}"


def test_decompiler_extracts_ellipse_and_layout_bg(tmp_path):
    from PIL import Image
    import io

    # Create test presentation
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    # Add a decorative shape to layout
    import copy
    temp_s = prs.slides.add_slide(layout)
    bg_shp = temp_s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4), Inches(7.5))
    layout.shapes._spTree.append(copy.deepcopy(bg_shp._element))

    # Drop temp slide
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

    # Slide 1 with an ellipse picture
    slide = prs.slides.add_slide(layout)
    im = Image.new("RGB", (100, 100), color=(100, 150, 200))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    pic = slide.shapes.add_picture(buf, Inches(1), Inches(1), Inches(3), Inches(3))
    for p in pic._element.xpath('.//*[local-name()="prstGeom"]'):
        p.set("prst", "ellipse")

    deck_path = tmp_path / "test_geom.pptx"
    prs.save(str(deck_path))

    decompiler = SlideDecompiler(str(deck_path), raw_mode=True)
    slides = decompiler.parse_deck()

    assert len(slides) == 1
    nodes = slides[0]["nodes"]

    # Verify is_layout_bg extracted
    layout_bg_nodes = [n for n in nodes if n.get("is_layout_bg")]
    assert len(layout_bg_nodes) >= 1, "Expected at least 1 layout background shape"

    # Verify ellipse geometry extracted
    ellipse_nodes = [n for n in nodes if n.get("geom_type") == "ellipse"]
    assert len(ellipse_nodes) == 1, "Expected 1 ellipse picture node"
    assert ellipse_nodes[0]["type"] == "PICTURE"


def test_rich_sample_deck_fidelity():
    deck_path = Path("backend/samples/sample_12slide_archetypes.pptx")
    assert deck_path.exists(), "Sample deck does not exist"

    decompiler = SlideDecompiler(str(deck_path), raw_mode=True)
    slides = decompiler.parse_deck()

    assert len(slides) == 12, "Sample deck should have 12 slides"

    # Ellipses on Slides 4, 7, 11
    s4_ellipses = [n for n in slides[3]["nodes"] if n.get("geom_type") == "ellipse"]
    s7_ellipses = [n for n in slides[6]["nodes"] if n.get("geom_type") == "ellipse"]
    s11_ellipses = [n for n in slides[10]["nodes"] if n.get("geom_type") == "ellipse"]

    assert len(s4_ellipses) == 1, "Slide 4 should have 1 circular hero picture"
    assert len(s7_ellipses) == 4, "Slide 7 should have 4 circular team headshots"
    assert len(s11_ellipses) == 3, "Slide 11 should have 3 circular summary pictures"

    # Layout background shapes present on all slides
    for s in slides:
        bg_nodes = [n for n in s["nodes"] if n.get("is_layout_bg")]
        assert len(bg_nodes) >= 1, f"Slide {s['slide_id']} missing layout background node"
