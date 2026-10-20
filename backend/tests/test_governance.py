"""Test strict corporate typography contract, contrast, and branding."""

import io
import os
import tempfile
import pytest
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from fastapi.testclient import TestClient

from brandmorph.brand.profile import BrandDNA
from brandmorph.engine.pipeline import MorphPipeline, MorphOptions
from decompiler import SlideDecompiler
from main import app


def test_corporate_typography_and_contrast_contract(tmp_path):
    client = TestClient(app)
    sample_path = "backend/samples/sample_legacy_pitch.pptx"
    assert os.path.exists(sample_path)

    with open(sample_path, "rb") as f:
        file_bytes = f.read()

    # 1. Test POST /api/morph with enterprise_navy
    res = client.post(
        "/api/morph",
        files={"file": ("sample.pptx", file_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"brand": "enterprise_navy"}
    )
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["status"] == "success"

    # Decode and inspect morphed presentation
    import base64
    morphed_bytes = base64.b64decode(data["deck_base64"])
    out_file = tmp_path / "morphed.pptx"
    out_file.write_bytes(morphed_bytes)

    prs = pptx.Presentation(str(out_file))

    title_texts = {"FY25 Operations Review", "Service Performance Summary", "Quarterly Trend"}
    found_titles = set()

    for slide_idx, slide in enumerate(prs.slides):
        # Assert slide has brand stamp and bottom framing
        shape_names = [s.name for s in slide.shapes]
        assert "BrandMorph_Logo_Stamp" in shape_names or "BrandMorphStamp" in shape_names, f"Slide {slide_idx} missing BrandMorph logo stamp"
        assert "BrandMorphBottomFrame" in shape_names, f"Slide {slide_idx} missing BrandMorphBottomFrame"

        for shape in slide.shapes:
            if shape.name in ("BrandMorph_Logo_Stamp", "BrandMorphStamp"):
                continue

            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    p_txt = p.text.strip()
                    is_title = p_txt in title_texts

                    for r in p.runs:
                        if not r.text.strip():
                            continue

                        # 1. Title Runs
                        if is_title:
                            found_titles.add(p_txt)
                            assert r.font.name == "Times New Roman", (
                                f"Slide {slide_idx} Title '{r.text}' font is {r.font.name}, expected Times New Roman"
                            )
                            assert r.font.bold is False, (
                                f"Slide {slide_idx} Title '{r.text}' bold is {r.font.bold}, expected False"
                            )
                            if r.font.size:
                                pt = r.font.size.pt
                                assert 24.0 <= pt <= 32.0, (
                                    f"Slide {slide_idx} Title '{r.text}' size {pt}pt not in [24, 32]"
                                )
                            assert r.font.color and r.font.color.type == 1, (
                                f"Slide {slide_idx} Title '{r.text}' missing explicit color"
                            )
                        # 2. Body Runs (non-title, non-subheading)
                        elif r.font.size and r.font.size.pt < 13.0:
                            assert r.font.name == "Arial", (
                                f"Slide {slide_idx} Body '{r.text}' font is {r.font.name}, expected Arial"
                            )
                            assert r.font.bold is False, (
                                f"Slide {slide_idx} Body '{r.text}' bold is {r.font.bold}, expected False"
                            )
                            assert r.font.italic is False, (
                                f"Slide {slide_idx} Body '{r.text}' italic is {r.font.italic}, expected False"
                            )
                            assert r.font.size.pt >= 8.0, (
                                f"Slide {slide_idx} Body '{r.text}' size {r.font.size.pt} < 8pt"
                            )
                            assert r.font.color and r.font.color.type == 1, (
                                f"Slide {slide_idx} Body '{r.text}' missing explicit color (inherited black bug)"
                            )

    assert found_titles == title_texts, f"Expected to find all titles: {title_texts}, found: {found_titles}"


def test_bright_shape_contrast_and_image_extraction(tmp_path):
    """Test bright shapes get dark text and image shapes emit PICTURE with data URIs."""
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Add bright cyan shape (#00FFFF) with text
    cyan_card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2)
    )
    cyan_card.fill.solid()
    cyan_card.fill.fore_color.rgb = RGBColor(0, 255, 255)  # Cyan (Luminance > 0.35)
    tf = cyan_card.text_frame
    p = tf.paragraphs[0]
    p.text = "Timeline Milestone 1"
    r = p.runs[0]
    r.font.size = Pt(14)

    # 2. Add an image shape with 1x1 test PNG
    # Minimal 1x1 valid PNG blob
    import base64
    tiny_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    img_stream = io.BytesIO(tiny_png)
    slide.shapes.add_picture(img_stream, Inches(6), Inches(1), Inches(3), Inches(2))

    test_deck = tmp_path / "test_contrast.pptx"
    prs.save(str(test_deck))

    # Run MorphPipeline
    dna = BrandDNA.load("enterprise_navy")
    morphed_deck = tmp_path / "morphed_contrast.pptx"
    MorphPipeline(dna, None, MorphOptions()).run(str(test_deck), str(morphed_deck))

    # Inspect morphed deck: text inside cyan card MUST have dark text (#030C1E)
    morphed_prs = pptx.Presentation(str(morphed_deck))
    m_slide = morphed_prs.slides[0]
    cyan_shape = [s for s in m_slide.shapes if s.has_text_frame and "Timeline" in s.text_frame.text][0]
    p = cyan_shape.text_frame.paragraphs[0]
    assert len(p.runs) > 0
    text_color = str(p.runs[0].font.color.rgb)
    assert text_color in ("030C1E", "000000"), f"Expected dark text on Cyan, got {text_color}"

    # Decompile and assert image was extracted as PICTURE
    dec = SlideDecompiler(str(morphed_deck))
    ast = dec.parse_deck()
    nodes = ast[0]["nodes"]
    pic_nodes = [n for n in nodes if n["type"] == "PICTURE"]
    assert len(pic_nodes) == 1, f"Expected 1 PICTURE node, found {len(pic_nodes)}"
    assert pic_nodes[0]["image_src"].startswith("data:image/")
    assert len(pic_nodes[0]["image_src"]) > 20
