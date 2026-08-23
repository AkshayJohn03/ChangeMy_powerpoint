"""Phase 12.6 Mutation Truth Layer."""
from __future__ import annotations

from dataclasses import dataclass, field

from pptx.util import Pt
from pptx.dml.color import RGBColor

from .brand_profile import BrandProfile


@dataclass
class MutationResult:
    status: str
    action: str
    shape_id: int = 0
    shape_name: str = ""
    changes: list = field(default_factory=list)
    reason: str = ""

    def to_dict(self):
        return {"status": self.status, "action": self.action,
                "shape_id": self.shape_id, "shape_name": self.shape_name,
                "changes": self.changes, "reason": self.reason}


def _lum(rgb):
    r, g, b = rgb[0] / 255, rgb[1] / 255, rgb[2] / 255
    return 0.299 * r + 0.587 * g + 0.114 * b


def _hex_str(rgb):
    if rgb is None:
        return "(theme)"
    return "#{:02X}{:02X}{:02X}".format(int(rgb[0]), int(rgb[1]), int(rgb[2]))


def _capture_run_state(run):
    s = {}
    try:
        s["font"] = run.font.name or "(inherited)"
    except Exception:
        s["font"] = "?"
    try:
        s["size_pt"] = round(run.font.size.pt, 1) if run.font.size else None
    except Exception:
        s["size_pt"] = None
    try:
        c = run.font.color.rgb
        s["color"] = str(c) if c else "(theme)"
    except Exception:
        s["color"] = "(theme)"
    return s


class MutationEngine:
    def __init__(self, profile):
        self.profile = profile

    def apply(self, shape, decision):
        action = decision.get("action", "")
        if action == "UNCHANGED":
            sid = getattr(shape, "shape_id", 0)
            return MutationResult("SKIPPED_NOOP", action, sid, shape.name)

        if action == "REMOVE":
            return self._remove(shape)

        if not shape.has_text_frame:
            sid = getattr(shape, "shape_id", 0)
            return MutationResult("UNSUPPORTED", action, sid, shape.name, reason="no text frame")

        before_list = []
        after_list = []
        changed_any = False
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                before = _capture_run_state(r)
                self._mutate_run(r, action)
                after = _capture_run_state(r)
                if before != after:
                    changed_any = True
                    before_list.append(before)
                    after_list.append(after)

        status = "CHANGED" if changed_any else "NO_CHANGE"
        changes = []
        if changed_any and before_list:
            b0, a0 = before_list[0], after_list[0]
            for k in b0:
                if b0[k] != a0[k]:
                    changes.append({"property": k, "before": b0[k], "after": a0[k]})

        reason = "{} applied with {} change(s)".format(action, len(changes)) if changed_any else "{} produced no visible change".format(action)
        sid = shape.shape_id
        return MutationResult(status, action, sid, shape.name, changes, reason)

    def _remove(self, shape):
        try:
            parent = shape._element.getparent()
            parent.remove(shape._element)
            return MutationResult("CHANGED", "REMOVE", shape.shape_id, shape.name, reason="element removed")
        except Exception as e:
            return MutationResult("FAILED", "REMOVE", shape.shape_id, shape.name, reason=str(e))

    def _mutate_run(self, run, action):
        bp = self.profile
        font = run.font

        if action == "RESTYLE":
            current = (font.name or "").lower()
            target = bp.body_font.lower()
            if current != target:
                font.name = bp.body_font

        elif action == "RESIZE":
            if font.size:
                pt_val = round(font.size.pt, 1)
                tier_target = self._resolve_size_tier(pt_val)
                if tier_target is not None and abs(pt_val - tier_target) >= 0.5:
                    font.size = Pt(tier_target)

        elif action == "RECOLOR":
            self._recolor_run(font)

    def _resolve_size_tier(self, pt_val):
        tiers = sorted(self.profile.font_tiers, key=lambda t: t["min"], reverse=True)
        for tier in tiers:
            if pt_val >= tier["min"]:
                return tier["target"]
        return None

    def _recolor_run(self, font):
        try:
            rgb = font.color.rgb
            if rgb is None:
                return
            lum = _lum(rgb)
            src_hex = _hex_str(rgb).upper().lstrip("#")
            legacy_variants = getattr(self.profile, "legacy_gold_variants", ())
            for cand in legacy_variants:
                cand_hex = "{:02X}{:02X}{:02X}".format(*cand)
                if src_hex == cand_hex:
                    font.color.rgb = self.profile.gold_rgb
                    return
            if lum < 0.25:
                font.color.rgb = self.profile.white_rgb
            elif 0.35 < lum < 0.55:
                font.color.rgb = self.profile.gold_rgb
        except Exception:
            pass
