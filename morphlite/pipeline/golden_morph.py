"""GoldenMorph — golden-layout-driven renderer (GLaD-Morph).

Design (from the agreed spec):
  * retrieval: anonymized data/metadata/golden_layout_index.json, ranked by
    archetype match + capacity fit + signature cosine;
  * Hungarian assignment (scipy linear_sum_assignment) between raw text
    blocks and the retrieved layout's slots -> golden reading order;
  * rendering on cloned master slides (lite template) with strict ECMA-376
    spPr child ordering (insert_spPr_child) and true native a:custGeom
    vector icons (6 line-art icons, 2400x2400 viewbox);
  * split layout: bullets left glass card + right connected step boxes
    with directional arrows; card grid + bullets card fallbacks.

Every glass card/step/badge is a native pptx shape (roundRect/oval) so the
output stays 100% editable and opens without repair warnings.
"""
from __future__ import annotations

import hashlib
import json
import math
import uuid
from pathlib import Path

import numpy as np
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt
from scipy.optimize import linear_sum_assignment

from ..training.build_golden_index import REPO, classify
from .template_filler import _duplicate_slide, _remove_first_n_slides
from .template_tag import (
    DEFAULT_TEMPLATE_PATH,
    TAG_RE,
    _purge_unreplaced_tags_in_shape,
    _replace_tag_in_shape,
    select_template_slide,
)

INDEX_PATH = REPO / "data" / "metadata" / "golden_layout_index.json"
EMU = 914400.0

# ECMA-376 CT_ShapeProperties child order (strict; never append blindly)
SPPR_ORDER = [
    "xfrm", "custGeom", "prstGeom", "noFill", "solidFill", "gradFill",
    "blipFill", "pattFill", "grpFill", "ln", "effectLst", "effectDag",
    "scene3d", "sp3d", "extLst",
]

GOLD = "FFB81C"
CARD = "081E36"
CARD_ALT = "0A2A4A"
WHITE = "FFFFFF"
MUTED = "B0C4DE"
FONT = "Arial"


def insert_spPr_child(spPr, child, order=SPPR_ORDER):
    """Insert child into spPr respecting the strict schema order."""
    tag = child.tag.split("}")[-1]
    idx = order.index(tag)
    for existing in spPr:
        etag = existing.tag.split("}")[-1]
        if etag in order and order.index(etag) > idx:
            existing.addprevious(child)
            return child
    spPr.append(child)
    return child


def _ring(cx, cy, r):
    k = 0.5523 * r
    return [
        ("M", cx, cy - r),
        ("C", cx + k, cy - r, cx + r, cy - k, cx + r, cy),
        ("C", cx + r, cy + k, cx + k, cy + r, cx, cy + r),
        ("C", cx - k, cy + r, cx - r, cy + k, cx - r, cy),
        ("C", cx - r, cy - k, cx - k, cy - r, cx, cy - r),
        ("Z",),
    ]


def _build_icons() -> dict:
    icons = {}
    icons["sensor_radar"] = (
        _ring(1200, 1200, 1000) + _ring(1200, 1200, 660) + _ring(1200, 1200, 320)
        + [("M", 1200, 1200), ("L", 1913, 487)] + _ring(1200, 1200, 70)
    )
    ops = [
        ("M", 500, 500), ("L", 1900, 500), ("L", 1900, 1900), ("L", 500, 1900), ("Z",),
        ("M", 900, 900), ("L", 1500, 900), ("L", 1500, 1500), ("L", 900, 1500), ("Z",),
    ]
    for x in (900, 1200, 1500):
        ops += [("M", x, 500), ("L", x, 240)]
        ops += [("M", x, 1900), ("L", x, 2160)]
    for y in (900, 1200, 1500):
        ops += [("M", 500, y), ("L", 240, y)]
        ops += [("M", 1900, y), ("L", 2160, y)]
    icons["cpu_chip"] = ops
    icons["shield_check"] = [
        ("M", 1200, 180), ("L", 2020, 560), ("L", 2020, 1280),
        ("L", 1200, 2220), ("L", 380, 1280), ("L", 380, 560), ("Z",),
        ("M", 800, 1160), ("L", 1120, 1480), ("L", 1680, 840),
    ]
    gear = _ring(1200, 1200, 560) + _ring(1200, 1200, 190)
    for i in range(8):
        a = math.radians(i * 45)
        gear += [
            ("M", 1200 + 560 * math.cos(a), 1200 + 560 * math.sin(a)),
            ("L", 1200 + 830 * math.cos(a), 1200 + 830 * math.sin(a)),
        ]
    icons["workflow_gear"] = gear
    icons["metric_trend"] = [
        ("M", 380, 300), ("L", 380, 2040), ("L", 2120, 2040),
        ("M", 560, 1720), ("L", 980, 1300), ("L", 1400, 1560), ("L", 2000, 700),
        ("M", 2000, 700), ("L", 1856, 760), ("M", 2000, 700), ("L", 1958, 852),
    ]
    icons["alert_triangle"] = (
        [("M", 1200, 260), ("L", 2180, 2040), ("L", 220, 2040), ("Z",),
         ("M", 1200, 860), ("L", 1200, 1500)]
        + _ring(1200, 1740, 80)
    )
    return icons


ICONS = _build_icons()

ICON_KEYWORDS = [
    ("sensor_radar", ("sensor", "radar", "lidar", "camera", "perception", "fusa")),
    ("cpu_chip", ("ecu", "chip", "processor", "compute", "hardware", "controller", "platform")),
    ("shield_check", ("safety", "compliance", "quality", "iso", "security", "validation", "homolog")),
    ("metric_trend", ("kpi", "metric", "performance", "growth", "target", "benchmark", "efficiency", "reduction")),
    ("alert_triangle", ("risk", "alert", "warning", "issue", "failure", "hazard", "challenge")),
    ("workflow_gear", ("process", "workflow", "pipeline", "development", "architecture", "system", "integration", "lifecycle")),
]


def _pick_icon(text: str) -> str:
    low = text.lower()
    for icon, kws in ICON_KEYWORDS:
        if any(k in low for k in kws):
            return icon
    return "workflow_gear"


def _custgeom_xml(icon: str):
    parts = []
    for op in ICONS[icon]:
        if op[0] == "M":
            parts.append(f'<a:moveTo><a:pt x="{int(round(op[1]))}" y="{int(round(op[2]))}"/></a:moveTo>')
        elif op[0] == "L":
            parts.append(f'<a:lnTo><a:pt x="{int(round(op[1]))}" y="{int(round(op[2]))}"/></a:lnTo>')
        elif op[0] == "C":
            parts.append(
                "<a:cubicBezTo>"
                f'<a:pt x="{int(round(op[1]))}" y="{int(round(op[2]))}"/>'
                f'<a:pt x="{int(round(op[3]))}" y="{int(round(op[4]))}"/>'
                f'<a:pt x="{int(round(op[5]))}" y="{int(round(op[6]))}"/>'
                "</a:cubicBezTo>"
            )
        elif op[0] == "Z":
            parts.append("<a:close/>")
    xml = (
        f'<a:custGeom {nsdecls("a")}>'
        "<a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>"
        '<a:rect l="l" t="t" r="r" b="b"/>'
        '<a:pathLst><a:path w="2400" h="2400" fill="none" stroke="1">'
        + "".join(parts)
        + "</a:path></a:pathLst></a:custGeom>"
    )
    return parse_xml(xml)


def _cosine(a, b) -> float:
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    na, nb = np.linalg.norm(a), np.linalg.norm(b)
    if na == 0 or nb == 0:
        return 0.0
    return float(np.dot(a, b) / (na * nb))


def _raw_signature(blocks) -> list:
    n = len(blocks)
    n_short = sum(1 for b in blocks if len(b) <= 38)
    n_long = n - n_short
    return [
        n / 20.0, n_short / 20.0, n_long / 20.0, 0.0, 0.0, 0.0, 0.0,
        1.0 if blocks else 0.0, n_short / max(n, 1), n_long / max(n, 1),
        0.0, min(n, 30) / 30.0,
    ]


class GoldenMorphPipeline:
    """Retrieval + Hungarian alignment + custGeom renderer."""

    def __init__(self, template_path: str | None = None, index_path: str | None = None):
        self.template_path = template_path or DEFAULT_TEMPLATE_PATH
        self.index = json.loads(Path(index_path or INDEX_PATH).read_text(encoding="utf-8"))
        self.entries = self.index["entries"]

    # -- retrieval -----------------------------------------------------------
    def retrieve(self, archetype: str, blocks, top_k: int = 3):
        pool = [e for e in self.entries if e["archetype"] == archetype]
        if not pool:
            pool = [e for e in self.entries if not e["archetype"].startswith("master_")]
        sig = _raw_signature(blocks)
        scored = []
        for e in pool:
            cos = _cosine(sig, e["signature"])
            cap = max(0, len(blocks) - e["n_text_slots"]) / 10.0
            scored.append((round(cos - 0.15 * cap, 4), e["id"]))
        scored.sort(key=lambda s: -s[0])
        ids = [s[1] for s in scored[:top_k]]
        by_id = {e["id"]: e for e in self.entries}
        return [by_id[i] for i in ids], scored[:top_k]

    def _align_order(self, texts, entry) -> list:
        if entry is None or not texts:
            return list(texts)
        slots = [s for s in entry["slots"] if s["role"] != "title"]
        if not slots:
            return list(texts)
        C = np.zeros((len(texts), len(slots)))
        for i, t in enumerate(texts):
            L = len(t)
            for j, s in enumerate(slots):
                cap = max(8.0, s["bbox_in"][2] * 55.0)
                cost = abs(L - min(L, cap)) / max(L, 1.0)
                if s["role"] == "bullet" and L <= 38:
                    cost += 0.15
                C[i, j] = cost
        ri, ci = linear_sum_assignment(C)
        pairs = sorted(zip(ri, ci), key=lambda p: (slots[p[1]]["bbox_in"][1], slots[p[1]]["bbox_in"][0]))
        order = [int(i) for i, _j in pairs]
        order += [i for i in range(len(texts)) if i not in order]
        return [texts[i] for i in order]

    # -- low-level drawing ----------------------------------------------------
    @staticmethod
    def _set_alpha(shp, val: int):
        spPr = shp._element.spPr
        solid = spPr.find(qn("a:solidFill"))
        if solid is None:
            return
        clr = solid.find(qn("a:srgbClr"))
        if clr is None:
            return
        alpha = parse_xml(f'<a:alpha {nsdecls("a")} val="{val}"/>')
        clr.append(alpha)

    def _add_glass_card(self, slide, x, y, w, h, fill=CARD, alpha=85000, border=GOLD, bw=1.0):
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shp.adjustments[0] = 0.065
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        self._set_alpha(shp, alpha)
        shp.line.color.rgb = RGBColor.from_string(border)
        shp.line.width = Pt(bw)
        shp.shadow.inherit = False
        return shp

    def _add_badge(self, slide, icon, x, y, size):
        b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        b.fill.solid()
        b.fill.fore_color.rgb = RGBColor.from_string(CARD_ALT)
        self._set_alpha(b, 60000)
        b.line.color.rgb = RGBColor.from_string(GOLD)
        b.line.width = Pt(0.75)
        b.shadow.inherit = False

        inset = size * 0.18
        ic = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x + inset), Inches(y + inset),
            Inches(size - 2 * inset), Inches(size - 2 * inset),
        )
        spPr = ic._element.spPr
        prst = spPr.find(qn("a:prstGeom"))
        if prst is not None:
            spPr.remove(prst)
        insert_spPr_child(spPr, _custgeom_xml(icon))
        ic.fill.background()
        ic.line.color.rgb = RGBColor.from_string(GOLD)
        ic.line.width = Pt(1.4)
        ic.shadow.inherit = False
        return ic

    def _add_text_lines(self, slide, lines, x, y, w, h, size=14, bullet=True, color=WHITE):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = ("•  " + line) if bullet else line
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.color.rgb = RGBColor.from_string(color)
        return tb

    @staticmethod
    def _delete_holders(slide, needles=("{bullets_text}", "{bullets}", "{subtitle}")):
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            txt = shape.text_frame.text
            if any(nd in txt for nd in needles):
                shape._element.getparent().remove(shape._element)

    def _fill_tags(self, slide, payload):
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            for m in TAG_RE.finditer(shape.text_frame.text):
                prefix, tag = m.groups()
                if prefix or tag not in payload:
                    continue
                _replace_tag_in_shape(shape, tag, str(payload[tag]))

    # -- layout renderers ------------------------------------------------------
    def _render_bullets_card(self, slide, W, H, bullets, title):
        x0, y0 = 0.65, 1.6
        x1, y1 = W - 0.65, H - 0.65
        self._add_glass_card(slide, x0, y0, x1 - x0, y1 - y0)
        icon = _pick_icon(title + " " + " ".join(bullets[:2]))
        self._add_badge(slide, icon, x0 + 0.32, y0 + 0.32, 0.7)
        size = 14 if len(bullets) <= 7 else (12 if len(bullets) <= 11 else 11)
        self._add_text_lines(slide, bullets, x0 + 0.34, y0 + 1.25, (x1 - x0) - 0.68, (y1 - y0) - 1.55, size)
        return [icon]

    def _render_split(self, slide, W, H, longs, steps, title):
        x0, y0 = 0.65, 1.6
        x1, y1 = W - 0.65, H - 0.65
        total_h = y1 - y0
        left_w = round((x1 - x0) * 0.45, 2)
        gap = 0.45
        right_x = x0 + left_w + gap
        right_w = x1 - right_x

        icon = _pick_icon(title + " " + (longs[0] if longs else ""))
        self._add_glass_card(slide, x0, y0, left_w, total_h)
        self._add_badge(slide, icon, x0 + 0.3, y0 + 0.3, 0.66)
        size = 14 if len(longs) <= 5 else 12
        if longs:
            self._add_text_lines(slide, longs, x0 + 0.32, y0 + 1.18, left_w - 0.64, total_h - 1.5, size)

        n = len(steps)
        arrow_gap = 0.3 if n > 1 else 0.0
        box_h = (total_h - (n - 1) * arrow_gap) / n
        while box_h < 0.55 and n > 2:
            n -= 1
            box_h = (total_h - (n - 1) * arrow_gap) / n
        steps = steps[:n]

        for i, t in enumerate(steps):
            y = y0 + i * (box_h + arrow_gap)
            self._add_glass_card(slide, right_x, y, right_w, box_h, fill=CARD_ALT, bw=0.75)
            tb = slide.shapes.add_textbox(
                Inches(right_x + 0.22), Inches(y + 0.04),
                Inches(right_w - 0.44), Inches(box_h - 0.08),
            )
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            num = p.add_run()
            num.text = f"{i + 1}.  "
            num.font.size = Pt(12)
            num.font.name = FONT
            num.font.bold = True
            num.font.color.rgb = RGBColor.from_string(GOLD)
            txt = p.add_run()
            txt.text = t
            txt.font.size = Pt(11 if len(t) > 46 else 12)
            txt.font.name = FONT
            txt.font.color.rgb = RGBColor.from_string(WHITE)
            if i < n - 1:
                arr = slide.shapes.add_shape(
                    MSO_SHAPE.DOWN_ARROW,
                    Inches(right_x + right_w / 2 - 0.11),
                    Inches(y + box_h + (arrow_gap - 0.22) / 2),
                    Inches(0.22), Inches(0.22),
                )
                arr.fill.solid()
                arr.fill.fore_color.rgb = RGBColor.from_string(GOLD)
                arr.line.fill.background()
                arr.shadow.inherit = False
        return [icon]

    # -- main ------------------------------------------------------------------
    def process(self, raw_path: str, output_path: str | None = None) -> dict:
        from ..models.content_model import extract_deck

        deck = extract_deck(raw_path)
        prs = pptx.Presentation(self.template_path)
        template_len = len(prs.slides)
        W = (prs.slide_width or 0) / EMU
        H = (prs.slide_height or 0) / EMU

        chosen = [
            select_template_slide(s, i, len(deck.slides), template_len)
            for i, s in enumerate(deck.slides)
        ]
        for t in chosen:
            _duplicate_slide(prs, t)
        _remove_first_n_slides(prs, template_len)

        closing = min(9, template_len - 1)
        diag = []
        for si, s in enumerate(deck.slides):
            slide = prs.slides[si]
            all_blocks = [b.text.strip() for b in s.text_blocks if b.text.strip()]
            blocks = [t for t in all_blocks if not (len(t) <= 2 and t.isdigit())] or all_blocks
            title = blocks[0] if blocks else ""
            body = blocks[1:]
            tpl_idx = chosen[si]
            entry = {"slide": si + 1, "template_slide": tpl_idx, "mode": "tag"}

            if si == 0 or tpl_idx == closing:
                if tpl_idx == closing and "thank" not in title.lower():
                    title = "Thank You"
                self._fill_tags(slide, {"title": title})
                diag.append(entry)
                continue

            n_short = sum(1 for b in blocks if len(b) <= 38)
            arch = classify({
                "n_images": 0,
                "n_text_slots": len(blocks),
                "n_card_slots": n_short,
                "n_bullets": len(blocks) - n_short,
                "has_flow_arrows": False,
            })
            entry["archetype"] = arch
            top_entries, top_scores = self.retrieve(arch, blocks)
            entry["retrieved"] = [{"id": i, "score": sc} for sc, i in top_scores]
            ordered = self._align_order(body, top_entries[0] if top_entries else None)
            steps = [t for t in ordered if len(t) <= 38]
            longs = [t for t in ordered if len(t) > 38]

            if tpl_idx == 2 and body:
                entry["mode"] = "card_grid"
                payload = {"title": title}
                head = steps[:15] if len(steps) > 16 else steps[:16]
                for i, t in enumerate(head, start=1):
                    payload[f"card_{i}"] = t
                if len(steps) > 16:
                    payload["card_16"] = " • ".join(steps[15:])
                self._fill_tags(slide, payload)
            elif tpl_idx == 3 and body:
                self._delete_holders(slide)
                if len(steps) >= 3 and longs:
                    entry["mode"] = "split"
                    entry["n_steps"] = len(steps[:5])
                    entry["icons"] = self._render_split(slide, W, H, longs, steps[:5], title)
                else:
                    entry["mode"] = "bullets_card"
                    entry["icons"] = self._render_bullets_card(slide, W, H, ordered, title)
            elif tpl_idx == 3:
                self._fill_tags(slide, {"title": title})
            else:
                joined = "\n".join(f"• {b}" for b in body)
                payload = {"title": title, "bullets_text": joined, "bullets": joined}
                if body:
                    payload["subtitle"] = body[0]
                for i, t in enumerate(body[:4], start=1):
                    payload[f"agenda_{i}"] = t
                self._fill_tags(slide, payload)
            diag.append(entry)

        for slide in prs.slides:
            for shape in slide.shapes:
                _purge_unreplaced_tags_in_shape(shape)

        out = output_path or f"_golden_out_{uuid.uuid4().hex[:6]}.pptx"
        prs.save(out)
        modes: dict[str, int] = {}
        for d in diag:
            modes[d["mode"]] = modes.get(d["mode"], 0) + 1
        return {
            "n_slides": len(prs.slides),
            "n_raw": len(deck.slides),
            "output": out,
            "template": self.template_path,
            "routing": [{"raw_slide": i + 1, "template_slide": t} for i, t in enumerate(chosen)],
            "slides": diag,
            "modes": modes,
            "hash": hashlib.sha256(Path(out).read_bytes()).hexdigest()[:12],
        }
