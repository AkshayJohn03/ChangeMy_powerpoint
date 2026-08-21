"""Training-grade PPTX AST extractor.

Produces one JSON deck file per source deck:
    data/extracted/ast/<bucket>/<batch>/<deck_id>.json

Design (PDR section 7):
  - persistent deck/slide/node ids
  - normalized 0..1 coordinates plus original inch values
  - full run-level typography and colors
  - group recursion with parent container links
  - table/chart detection
  - no base64 image payloads
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from ..ast.node_ids import deck_id, slide_id, node_id, slugify


def _inch(v) -> float | None:
    try:
        return round(v.inches, 4)
    except Exception:
        return None


def _norm(v: float | None, total: float | None) -> float | None:
    if v is None or not total:
        return None
    return round(max(0.0, min(1.0, v / total)), 5)


def _rgb(color_fmt) -> str | None:
    try:
        if color_fmt and color_fmt.type is not None and color_fmt.rgb is not None:
            return str(color_fmt.rgb)
    except Exception:
        pass
    return None


def _fill_color(shape) -> str | None:
    try:
        fill = shape.fill
        if fill.type == 1:
            return _rgb(fill.fore_color)
    except Exception:
        pass
    return None


def _alignment(p) -> str | None:
    try:
        if p.alignment is None:
            return None
        return {
            PP_ALIGN.LEFT: "LEFT",
            PP_ALIGN.CENTER: "CENTER",
            PP_ALIGN.RIGHT: "RIGHT",
            PP_ALIGN.JUSTIFY: "JUSTIFY",
        }.get(p.alignment, str(p.alignment))
    except Exception:
        return None


def _extract_runs(paragraph) -> list[dict]:
    runs = []
    for r in paragraph.runs:
        font = r.font
        runs.append(
            {
                "text": r.text,
                "font_family": font.name,
                "font_size_pt": round(font.size.pt, 2) if font.size else None,
                "bold": font.bold,
                "italic": font.italic,
                "underline": bool(font.underline) if font.underline is not None else None,
                "color": _rgb(font.color),
            }
        )
    return runs


def _text_shape_payload(shape) -> dict:
    paragraphs = []
    for p in shape.text_frame.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        run_sizes = [r["font_size_pt"] for r in _extract_runs(p) if r["font_size_pt"]]
        paragraphs.append(
            {
                "text": text,
                "level": p.level,
                "align": _alignment(p),
                "font_size_pt": run_sizes[0] if run_sizes else None,
                "runs": _extract_runs(p),
            }
        )
    return {"paragraphs": paragraphs}


def _shape_type(shape) -> str:
    st = shape.shape_type
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "TEXT"
    mapping = {
        MSO_SHAPE_TYPE.AUTO_SHAPE: "SHAPE",
        MSO_SHAPE_TYPE.FREEFORM: "SHAPE",
        MSO_SHAPE_TYPE.PICTURE: "IMAGE",
        MSO_SHAPE_TYPE.LINKED_PICTURE: "IMAGE",
        MSO_SHAPE_TYPE.LINE: "LINE",
        MSO_SHAPE_TYPE.GROUP: "GROUP",
        MSO_SHAPE_TYPE.TABLE: "TABLE",
        MSO_SHAPE_TYPE.CHART: "CHART",
        MSO_SHAPE_TYPE.TEXT_BOX: "TEXT",
    }
    result = mapping.get(st)
    if result == "SHAPE":
        try:
            name = str(shape.auto_shape_type)
            if "ARROW" in name or "CHEVRON" in name:
                return "ARROW"
        except Exception:
            pass
    return result or "OTHER"


def _table_payload(shape) -> dict | None:
    try:
        tbl = shape.table
        rows = []
        for row in tbl.rows:
            rows.append([c.text.strip() for c in row.cells])
        return {"n_rows": len(rows), "n_cols": len(rows[0]) if rows else 0, "cells": rows}
    except Exception:
        return None


def _extract_node(shape, sid: str, counter: list[int], parent_id: str | None, sw: float, sh: float, depth: int) -> dict | None:
    left, top = _inch(shape.left), _inch(shape.top)
    width, height = _inch(shape.width), _inch(shape.height)
    ntype = _shape_type(shape)

    node: dict[str, Any] = {
        "id": node_id(sid, counter[0]),
        "type": ntype,
        "shape_id": getattr(shape, "shape_id", None),
        "parent_container_id": parent_id,
        "container_depth": depth,
        "bbox_norm": {
            "x": _norm(left, sw),
            "y": _norm(top, sh),
            "w": _norm(width, sw),
            "h": _norm(height, sh),
        },
        "bbox_in": {"x": left, "y": top, "w": width, "h": height},
        "rotation": getattr(shape, "rotation", 0.0),
        "name": shape.name,
    }
    counter[0] += 1

    if ntype in ("TEXT", "SHAPE", "ARROW", "OTHER"):
        node["fill_color"] = _fill_color(shape)
        try:
            if shape.has_text_frame:
                payload = _text_shape_payload(shape)
                if payload["paragraphs"]:
                    node["paragraphs"] = payload["paragraphs"]
        except Exception:
            pass
    elif ntype == "TABLE":
        payload = _table_payload(shape)
        if payload:
            node["table"] = payload
    elif ntype == "CHART":
        try:
            node["chart_type"] = str(shape.chart.chart_type)
        except Exception:
            node["chart_type"] = None
    elif ntype == "IMAGE":
        try:
            blob = shape.image.blob
            node["image_ext"] = shape.image.ext
            node["image_md5_8"] = hashlib.md5(blob).hexdigest()[:8]
        except Exception:
            pass

    children = []
    if ntype == "GROUP":
        try:
            for child in shape.shapes:
                child_node = _extract_node(child, sid, counter, node["id"], sw, sh, depth + 1)
                if child_node:
                    children.append(child_node)
        except Exception:
            pass
    if children:
        node["children"] = children

    if ntype != "GROUP":
        has_content = bool(node.get("paragraphs")) or ntype in ("IMAGE", "TABLE", "CHART")
        if not has_content and width is not None and height is not None:
            if width * height < 0.02:
                return None
    return node


def _slide_background(slide) -> str | None:
    try:
        fill = slide.background.fill
        if fill.type == 1:
            return _rgb(fill.fore_color)
    except Exception:
        pass
    return None


def extract_deck(path: Path, bucket: str, batch: str, name_key: str, md5: str) -> dict:
    prs = pptx.Presentation(str(path))
    sw_in = _inch_val(prs.slide_width)
    sh_in = _inch_val(prs.slide_height)
    did = deck_id(bucket, batch, path.name)

    slides_out = []
    for idx, slide in enumerate(prs.slides):
        sid = slide_id(did, idx + 1)
        counter = [0]
        nodes = []
        for shape in slide.shapes:
            n = _extract_node(shape, sid, counter, None, sw_in, sh_in, 0)
            if n:
                nodes.append(n)
                if n["type"] == "GROUP":
                    pass
        slides_out.append(
            {
                "slide_id": sid,
                "index": idx + 1,
                "layout": slide.slide_layout.name,
                "background_color": _slide_background(slide),
                "nodes": nodes,
            }
        )

    return {
        "schema_version": "1.0.0",
        "deck_id": did,
        "source": {"path": str(path), "bucket": bucket, "batch": batch, "md5": md5, "name_key": name_key},
        "dimensions_in": {"width": sw_in, "height": sh_in},
        "n_slides": len(slides_out),
        "slides": slides_out,
    }


def _inch_val(emu) -> float:
    try:
        return round(emu / 914400.0, 4)
    except Exception:
        return 13.333


def flatten_nodes(nodes: list[dict]) -> list[dict]:
    out = []
    for n in nodes:
        out.append(n)
        if n.get("children"):
            out.extend(flatten_nodes(n["children"]))
    return out


def node_text(node: dict) -> str:
    parts = [p["text"] for p in node.get("paragraphs") or []]
    if node.get("table"):
        for row in node["table"]["cells"]:
            parts.extend(t for t in row if t)
    return " ".join(parts)


def slide_text(deck_json: dict, slide_index: int) -> str:
    parts: list[str] = []
    for n in flatten_nodes(deck_json["slides"][slide_index]["nodes"]):
        t = node_text(n)
        if t:
            parts.append(t)
    return " ".join(parts)


def save_deck(deck: dict, out_root: Path) -> Path:
    src = deck["source"]
    out_dir = out_root / src["bucket"] / src["batch"]
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{deck['deck_id']}.json"
    out_path.write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
    return out_path
