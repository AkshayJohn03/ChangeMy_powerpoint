"""Verification gate 1: unit-test the inheritance resolver.

Run: python -m morphlite.tests.test_inheritance
"""
from __future__ import annotations

import sys
from pathlib import Path

import pptx

from ..ast.effective import DeckTypographyResolver

BRANDMORPH_ROOT = Path(__file__).resolve().parents[2]


def _any_corpus_deck() -> Path | None:
    for pattern in ("data/source/paired/raw/**/*.pptx", "data/source/branded_only/**/*.pptx"):
        hits = sorted((BRANDMORPH_ROOT / pattern).parent.glob(pattern.split("/")[-1])) if False else sorted(
            p for p in (BRANDMORPH_ROOT / pattern.split("/**/")[0]).rglob(pattern.split("/**/")[1])
        )
        if hits:
            return hits[0]
    return None


def build_synthetic(tmp_path: Path) -> Path:
    prs = pptx.Presentation()
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    title = slide.shapes.title
    title.text = "Inherited title text"

    box = slide.shapes.add_textbox(0, 0, 4, 1)
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Explicit 28pt Arial"
    run.font.size = pptx.util.Pt(28)
    run.font.name = "Arial"
    run.font.bold = True

    out = tmp_path / "synthetic.pptx"
    prs.save(str(out))
    return out


def test_synthetic():
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        path = build_synthetic(Path(td))
        with DeckTypographyResolver(str(path)) as r:
            slide0 = r.resolve_slide(0)
        shapes_by_text = {}
        prs = pptx.Presentation(str(path))
        for sp in prs.slides[0].shapes:
            if sp.has_text_frame and sp.text_frame.text.strip():
                shapes_by_text[str(sp.shape_id)] = sp.text_frame.text.strip()

        explicit = None
        inherited = None
        for sid, eff in slide0.items():
            fam = eff.get("font_family", {}).get("value", "")
            if "Explicit" in shapes_by_text.get(sid, ""):
                explicit = eff
            elif "Inherited title" in shapes_by_text.get(sid, ""):
                inherited = eff

        assert explicit is not None, "explicit textbox not found in resolution"
        assert explicit["font_size_pt"]["value"] == 28.0, explicit
        assert explicit["font_size_pt"]["origin"] == "run", explicit["font_size_pt"]
        assert explicit["font_family"]["value"].lower() == "arial", explicit["font_family"]
        assert explicit["bold"]["value"] is True

        assert inherited is not None, "title placeholder not found in resolution"
        assert inherited["font_family"]["value"], "family unresolved for placeholder"
        assert inherited["font_family"]["origin"] != "run"
        assert inherited["font_size_pt"] is not None and inherited["font_size_pt"]["value"], "size unresolved for placeholder"

        print("synthetic OK")
        print("  explicit:", {k: v for k, v in explicit.items() if not k.startswith('_')})
        print("  inherited:", {k: v for k, v in inherited.items() if not k.startswith('_')})


def test_corpus_smoke():
    corpus = _any_corpus_deck()
    assert corpus is not None, "no anonymized corpus decks found under data/source"
    with DeckTypographyResolver(str(corpus)) as r:
        origin_counter = {}
        total_shapes = 0
        size_resolved = family_resolved = 0
        for idx in range(len(r.slide_parts)):
            res = r.resolve_slide(idx)
            for sid, eff in res.items():
                total_shapes += 1
                s = eff.get("font_size_pt")
                f = eff.get("font_family")
                if s and s.get("value"):
                    size_resolved += 1
                    origin_counter[s["origin"]] = origin_counter.get(s["origin"], 0) + 1
                else:
                    origin_counter["size_unresolved"] = origin_counter.get("size_unresolved", 0) + 1
                if f and f.get("value"):
                    family_resolved += 1
                    origin_counter[f"fam:{f['origin']}"] = origin_counter.get(f"fam:{f['origin']}", 0) + 1

    size_rate = size_resolved / max(total_shapes, 1)
    family_rate = family_resolved / max(total_shapes, 1)
    print(f"corpus smoke: shapes={total_shapes} size_resolved={size_rate:.1%} family_resolved={family_rate:.1%}")
    print("origins:", dict(sorted(origin_counter.items(), key=lambda kv: -kv[1])))
    assert size_rate >= 0.80, f"effective size coverage too low: {size_rate:.2%}"
    assert family_rate >= 0.90, f"effective family coverage too low: {family_rate:.2%}"


if __name__ == "__main__":
    test_synthetic()
    test_corpus_smoke()
    print("GATE 1 PASSED")
    sys.exit(0)
