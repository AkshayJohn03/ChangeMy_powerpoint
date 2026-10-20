import posixpath
import re
import sys
import zipfile

sys.path.insert(0, r"C:\Users\Akshay.JOHN-XAVIER\.gemini\antigravity-ide\scratch\BrandMorph")
import pptx
from morphlite.pipeline.template_tag import TAG_RE

lite = r"C:\Users\Akshay.JOHN-XAVIER\.gemini\antigravity-ide\scratch\BrandMorph\data\brand_template\Akkodis_Sanitized_Master_Lite.pptx"
z = zipfile.ZipFile(lite)
print("zip test:", z.testzip() is None)

p = pptx.Presentation(lite)
print("slides", len(p.slides), "masters", len(p.slide_masters),
      "layouts", sum(len(m.slide_layouts) for m in p.slide_masters))

skip = ("#", "/", "$", "%", "~")
for i, s in enumerate(p.slides):
    tags = sorted({
        m.group(2)
        for sh in s.shapes if sh.has_text_frame
        for m in TAG_RE.finditer(sh.text_frame.text)
        if m.group(1) not in skip
    })
    print("slide", i, len(tags), tags[:8])

names = set(z.namelist())
bad = []
for n in names:
    if not n.endswith(".rels"):
        continue
    root = posixpath.dirname(posixpath.dirname(n))
    for rel in re.findall(r"<Relationship ([^>]+)>", z.read(n).decode("utf-8")):
        if "External" in rel:
            continue
        t = re.search(r'Target="([^"]+)"', rel)
        if not t:
            continue
        target = posixpath.normpath(posixpath.join(root, t.group(1)))
        if target not in names:
            bad.append((n, t.group(1)))
print("bad rels:", len(bad), bad[:5])
