import os
import sys
import time

REPO = r"C:\Users\Akshay.JOHN-XAVIER\.gemini\antigravity-ide\scratch\BrandMorph"
sys.path.insert(0, REPO)
os.chdir(REPO)

import pptx
from morphlite.pipeline.template_tag import DEFAULT_TEMPLATE_PATH, TemplateTagPipeline

print("template:", DEFAULT_TEMPLATE_PATH)
p = TemplateTagPipeline()
print("tpl slides:", len(p.tpl.slides))

out = os.path.join(REPO, "scripts", "_verify_tag_out.pptx")
t0 = time.time()
res = p.process(
    r"C:\Users\Akshay.JOHN-XAVIER\OneDrive - Akkodis\Documents\Me\TestPPT.pptx", out
)
print("secs:", round(time.time() - t0, 1))
print("n_slides:", res["n_slides"], "n_raw:", res["n_raw"])
print("routing:", [(r["raw_slide"], r["template_slide"]) for r in res["routing"]])

prs = pptx.Presentation(out)
leftover = sum(
    1
    for s in prs.slides
    for sh in s.shapes
    if sh.has_text_frame and "{" in sh.text_frame.text and "}" in sh.text_frame.text
)
print("leftover tag shapes:", leftover)

last = prs.slides[-1]
texts = [sh.text_frame.text.strip() for sh in last.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
print("last slide texts:", texts[:6])
print("out MB:", round(os.path.getsize(out) / 1e6, 2))
