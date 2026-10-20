import os
import sys
import time
import zipfile

REPO = r"C:\Users\Akshay.JOHN-XAVIER\.gemini\antigravity-ide\scratch\BrandMorph"
sys.path.insert(0, REPO)
os.chdir(REPO)

import pptx
from lxml import etree

from morphlite.pipeline.golden_morph import GoldenMorphPipeline

out = os.path.join(REPO, "scripts", "_verify_golden_out.pptx")
pipe = GoldenMorphPipeline()
print("template:", pipe.template_path)
print("index entries:", len(pipe.entries))

t0 = time.time()
res = pipe.process(
    r"C:\Users\Akshay.JOHN-XAVIER\OneDrive - Akkodis\Documents\Me\TestPPT.pptx", out
)
print("secs:", round(time.time() - t0, 1))
print("n_slides:", res["n_slides"], "n_raw:", res["n_raw"])
print("modes:", res["modes"])
for d in res["slides"]:
    extra = ""
    if d["mode"] in ("split", "bullets_card"):
        extra = f" icons={d.get('icons')} retrieved={[r['id'] for r in d.get('retrieved', [])][:2]}"
    if d["mode"] in ("card_grid", "tag"):
        extra = f" retrieved={[r['id'] for r in d.get('retrieved', [])][:2]}" if "retrieved" in d else ""
    print(f"  slide {d['slide']:>2} -> tpl {d['template_slide']} [{d['mode']}]{extra}")

prs = pptx.Presentation(out)
leftover = sum(
    1
    for s in prs.slides
    for sh in s.shapes
    if sh.has_text_frame and "{" in sh.text_frame.text and "}" in sh.text_frame.text
)
print("leftover tag shapes:", leftover)

last_texts = [
    sh.text_frame.text.strip()
    for sh in prs.slides[-1].shapes
    if sh.has_text_frame and sh.text_frame.text.strip()
]
print("slide 29 texts:", last_texts[:5])

z = zipfile.ZipFile(out)
n_custgeom = 0
bad_xml = []
for name in z.namelist():
    if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
        raw = z.read(name)
        n_custgeom += raw.count(b"<a:custGeom")
        try:
            etree.fromstring(raw)
        except Exception as e:
            bad_xml.append((name, str(e)))
print("custGeom icon shapes:", n_custgeom)
print("bad xml:", bad_xml)
print("out MB:", round(os.path.getsize(out) / 1e6, 2))
